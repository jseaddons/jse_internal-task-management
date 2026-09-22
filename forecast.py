#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Forecast / loading plan — assign work to the team for the next X days.

Shares project lists (people, categories, zones, levels, models, tasks) with
the daily task app. Forecast rows are kept forever so a new X-day plan does
not wipe the last one. Comparison reads daily `tasks` (hours actually logged)
and never writes to that table.
"""
import html
import io
import json
import re
import sqlite3
import unicodedata
import datetime as dt
from itertools import product

SERVICES = ("Modelling Works", "Sheet Work")
SCALES = ("1:20", "1:25", "1:50", "1:75", "1:100", "1:200", "1:250", "1:500", "NTS")
PAPER_SIZES = ("A0", "A1", "A2", "A3", "A4")
HOURS_PER_DAY = 8.25
INK = "1F4E5F"


def ensure_public_holidays(con):
    con.execute("""
        CREATE TABLE IF NOT EXISTS public_holidays (
            day TEXT PRIMARY KEY,
            name TEXT
        )
    """)
    n = con.execute("SELECT COUNT(*) FROM public_holidays").fetchone()[0]
    if not n:
        con.execute(
            "INSERT OR IGNORE INTO public_holidays(day, name) VALUES (?,?)",
            ("2026-09-14", "Public holiday"))
        con.execute(
            "INSERT OR IGNORE INTO public_holidays(day, name) VALUES (?,?)",
            ("2026-10-02", "Gandhi Jayanti"))


def load_public_holiday_days(db_path=None, con=None):
    own = con is None
    if own:
        con = connect(db_path)
    ensure_public_holidays(con)
    days = {r[0] for r in con.execute("SELECT day FROM public_holidays")}
    if own:
        con.commit()
        con.close()
    return days


def is_weekly_off(day):
    """Sunday, plus the 2nd and 4th Saturday of the month. Other Saturdays are working days."""
    if isinstance(day, str):
        day = dt.date.fromisoformat(day[:10])
    if day.weekday() == 6:
        return True
    if day.weekday() == 5:
        return ((day.day - 1) // 7) + 1 in (2, 4)
    return False


def is_public_holiday(day, holidays=None):
    if isinstance(day, str):
        iso = day[:10]
    else:
        iso = day.isoformat()
    return bool(holidays) and iso in holidays


def is_leave_day(day, holidays=None):
    return is_weekly_off(day) or is_public_holiday(day, holidays)

STYLE = """<style>
  :root{--ink:#1a2230;--muted:#5b6675;--rule:#d5dbe3;--accent:#1f5a86;--ok:#1f7a4d;
    --bg:#f5f7fa;--card:#fff;--warn:#a33;--warnbg:#fde8e8}
  *{box-sizing:border-box} body{margin:0;background:var(--bg);color:var(--ink);
    font:15px/1.5 "Segoe UI",system-ui,Arial,sans-serif}
  .wrap{max-width:1180px;margin:0 auto;padding:24px 18px 60px}
  h1{font-size:20px;margin:0 0 4px} .sub{color:var(--muted);font-size:13px;margin:0 0 18px}
  .nav{margin-bottom:14px;font-size:13px} .nav a{color:var(--accent);text-decoration:none} .nav b{color:var(--ink)}
  .whoami{font-size:12px;color:var(--muted);margin:0 0 14px;padding:6px 10px;background:var(--card);
    border:1px solid var(--rule);border-radius:6px;display:inline-block} .whoami b{color:var(--ink)}
  .ok{background:#e6f4ec;border:1px solid var(--ok);color:var(--ok);padding:10px 14px;
    border-radius:8px;margin-bottom:16px;font-weight:600}
  .err{background:var(--warnbg);border:1px solid var(--warn);color:var(--warn);padding:10px 14px;
    border-radius:8px;margin-bottom:16px;font-weight:600}
  form.card, .card{background:var(--card);border:1px solid var(--rule);border-radius:10px;padding:16px 18px;margin-bottom:16px}
  .grid{display:grid;grid-template-columns:1fr 1fr;gap:14px 18px}
  .full{grid-column:1/-1}
  label{display:block;font-size:12px;letter-spacing:.02em;color:var(--muted);
    text-transform:uppercase;margin-bottom:4px;font-weight:600}
  input,select,textarea{width:100%;padding:9px 10px;border:1px solid var(--rule);
    border-radius:7px;font:inherit;background:#fff;color:var(--ink)}
  input:focus,select:focus,textarea:focus{outline:2px solid var(--accent);border-color:var(--accent)}
  .actions{margin-top:16px;display:flex;gap:10px;align-items:center;flex-wrap:wrap}
  button,.btn{background:var(--accent);color:#fff;border:0;border-radius:8px;padding:10px 18px;
    font-weight:600;font-size:14px;cursor:pointer;text-decoration:none;display:inline-block}
  button:hover,.btn:hover{filter:brightness(1.08)}
  button.grey,.btn.grey{background:#6b7280}
  button.green,.btn.green{background:#1f7a4d}
  button.danger{background:#a33;padding:6px 12px;font-size:13px}
  a.btn.sm, button.sm{padding:6px 12px;font-size:13px}
  td.act, th.act{white-space:nowrap;width:1%}
  .act form{display:inline}
  .hint{color:var(--muted);font-size:12px}
  h2{font-size:14px;margin:22px 0 8px}
  .scroll{overflow-x:auto;border:1px solid var(--rule);border-radius:10px;background:var(--card)}
  table{width:100%;border-collapse:collapse;font-size:13px}
  th,td{text-align:left;padding:7px 9px;border-bottom:1px solid var(--rule);white-space:nowrap}
  th{background:#eef2f6;color:var(--muted);text-transform:uppercase;font-size:11px;letter-spacing:.04em}
  td.num,th.num{text-align:right}
  td.over{background:var(--warnbg);color:var(--warn);font-weight:600}
  td.okh{background:#e6f4ec}
  td.short{background:#fff6e5;color:#8a5a00}
  td.empty{color:#c5ccd4}
  tr.total td{border-top:2px solid var(--rule)}
  .periods{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin:0 0 4px}
  .periods a,.periods .lab{border:1px solid var(--rule);border-radius:6px;padding:6px 10px;
    font-size:13px;text-decoration:none;color:var(--accent);background:#fff}
  .periods a.btn,.periods a.btn.green{background:#1f7a4d;color:#fff;border-color:#1f7a4d;
    font-weight:600;padding:8px 14px;font-size:14px}
  .periods .lab{color:var(--muted);border:0;padding-left:0}
  .periods .planchip{display:inline-flex;align-items:stretch;border:1px solid var(--rule);
    border-radius:6px;overflow:hidden;background:#fff}
  .periods .planchip a{border:0;border-radius:0}
  .periods .planchip form{margin:0;display:flex}
  .periods .planchip button.rm{background:#a33;padding:6px 8px;font-size:11px;
    letter-spacing:.04em;border-radius:0;min-width:auto}
  .cellbtn{all:unset;cursor:pointer;display:block;width:100%;text-align:right}
  .window{display:flex;flex-wrap:wrap;gap:10px;align-items:end}
  .window > div{min-width:140px}
  .exports{display:flex;gap:10px;flex-wrap:wrap;margin:8px 0 18px}
  .grid3{display:grid;grid-template-columns:1fr 1fr 1fr;gap:14px}
  .tickcol label .bulk, label .bulk{font-weight:400}
  .bulk{margin-left:8px;display:inline-flex;gap:6px;vertical-align:middle}
  .bulk button{background:none;border:1px solid var(--rule);color:var(--accent);
    border-radius:4px;padding:1px 8px;font-size:11px;cursor:pointer;text-transform:none;
    letter-spacing:0;font-weight:600}
  .bulk button:hover{background:var(--accent);color:#fff;border-color:var(--accent)}
  .ticks{max-height:210px;overflow-y:auto;border:1px solid var(--rule);border-radius:6px;
    padding:8px 10px;background:#fff;display:flex;flex-wrap:wrap;gap:6px 14px}
  .ticks .grp{width:100%;margin:6px 0 2px;font-size:11px;letter-spacing:.04em;
    text-transform:uppercase;color:var(--muted)}
  .ticks .grp:first-child{margin-top:0}
  .ticks .none{color:var(--muted);font-size:13px}
  .ticks label{display:flex;align-items:center;gap:6px;text-transform:none;font-size:13px;
    color:var(--ink);font-weight:400;margin:0;width:auto}
  .ticks input{width:auto;padding:0}
  .filter{margin:0 0 6px;padding:6px 8px;font-size:13px}
  .svc{display:flex;gap:8px;flex-wrap:wrap}
  .svc label{display:flex;align-items:center;gap:8px;text-transform:none;font-size:14px;
    font-weight:600;color:var(--ink);margin:0;border:1px solid var(--rule);border-radius:8px;
    padding:8px 14px;cursor:pointer;background:#fff}
  .svc input{width:auto;padding:0}
  .days{display:flex;flex-wrap:wrap;gap:8px}
  .days label{display:flex;align-items:center;gap:6px;text-transform:none;font-size:13px;
    font-weight:400;color:var(--ink);margin:0;border:1px solid var(--rule);border-radius:6px;
    padding:6px 10px;background:#fff;cursor:pointer}
  .days input{width:auto;padding:0}
  .countbar{background:#eef2f6;border-radius:8px;padding:10px 12px;font-size:13px;margin-top:12px}
  .countbar b{color:var(--accent)}
  .floordays{display:grid;grid-template-columns:repeat(auto-fill,minmax(170px,1fr));gap:8px;margin-top:6px}
  .frow{display:flex;align-items:center;gap:8px;border:1px solid var(--rule);border-radius:6px;
    padding:6px 8px;background:#fff}
  .frow .nm{flex:1;font-size:13px;font-weight:600}
  .frow input[type=number]{width:72px;padding:6px 8px}
  .prog{display:flex;height:22px;border-radius:7px;overflow:hidden;background:#e5e9ef;margin:10px 0 8px}
  .prog span{display:flex;align-items:center;justify-content:center;font-size:11px;font-weight:700;color:#fff;min-width:0}
  .prog .done{background:#1f7a4d} .prog .wip{background:#c47b12} .prog .todo{background:#94a3b8}
  .progstats{display:flex;flex-wrap:wrap;gap:14px 22px;font-size:13px;margin:0 0 10px}
  .progstats b{font-size:16px}
  .section{display:flex;flex-wrap:wrap;align-items:center;justify-content:space-between;
    gap:10px;margin:28px 0 10px}
  .section h2{margin:0}
  details.more{background:var(--card);border:1px solid var(--rule);border-radius:10px;
    padding:10px 16px;margin:0 0 16px}
  details.more > summary{cursor:pointer;font-weight:600;color:var(--accent);list-style:none}
  details.more > summary::-webkit-details-marker{display:none}
  details.more > summary::before{content:"+ ";opacity:.6}
  details.more[open] > summary::before{content:"− "}
  details.more .card{border:0;padding:12px 0 0;margin:0}
  .teamfold{margin-top:12px}
  .teamfold > summary{cursor:pointer;font-size:12px;font-weight:600;letter-spacing:.02em;
    text-transform:uppercase;color:var(--muted)}
  .steps{display:flex;flex-wrap:wrap;gap:8px;margin:0 0 14px}
  .steps span{background:#eef2f6;border-radius:999px;padding:4px 10px;font-size:12px;
    color:var(--muted)}
  .steps span b{color:var(--accent)}
  @media (max-width:900px){.grid3{grid-template-columns:1fr 1fr}}
  @media (max-width:640px){.grid3{grid-template-columns:1fr}}
</style>"""


def connect(db_path):
    con = sqlite3.connect(db_path, timeout=15)
    con.execute("PRAGMA foreign_keys = ON")
    con.execute("PRAGMA busy_timeout = 15000")
    con.row_factory = sqlite3.Row
    return con


def ensure_forecast_schema(db_path):
    con = connect(db_path)
    con.executescript("""
        CREATE TABLE IF NOT EXISTS forecast_assignments (
            id INTEGER PRIMARY KEY,
            work_date TEXT NOT NULL,
            project_code TEXT NOT NULL REFERENCES projects(code),
            service TEXT NOT NULL CHECK (service IN ('Modelling Works', 'Sheet Work')),
            person_id INTEGER NOT NULL REFERENCES people(id),
            category_id INTEGER REFERENCES categories(id),
            building_id INTEGER REFERENCES buildings(id),
            level_id INTEGER REFERENCES levels(id),
            area REAL,
            estimated_hours REAL NOT NULL,
            task_name_id INTEGER REFERENCES task_names(id),
            task_sub_task_id INTEGER REFERENCES task_sub_tasks(id),
            model_id INTEGER REFERENCES models(id),
            drawing_name TEXT,
            drawing_number TEXT,
            scale TEXT,
            paper_size TEXT,
            notes TEXT,
            entered_by INTEGER REFERENCES people(id)
        );
        CREATE INDEX IF NOT EXISTS ix_forecast_proj_date
            ON forecast_assignments(project_code, work_date);
        CREATE TABLE IF NOT EXISTS forecast_subtask_budget (
            id INTEGER PRIMARY KEY,
            project_code TEXT NOT NULL REFERENCES projects(code) ON DELETE CASCADE,
            category_id INTEGER REFERENCES categories(id),
            task_sub_task_id INTEGER NOT NULL REFERENCES task_sub_tasks(id),
            building_id INTEGER REFERENCES buildings(id),
            level_id INTEGER REFERENCES levels(id),
            hours REAL NOT NULL
        );
        CREATE UNIQUE INDEX IF NOT EXISTS ix_forecast_budget
            ON forecast_subtask_budget(
                project_code, task_sub_task_id,
                IFNULL(building_id, 0), IFNULL(level_id, 0));
    """)
    cols = {r[1] for r in con.execute("PRAGMA table_info(forecast_assignments)")}
    if "created_at" not in cols:
        con.execute("ALTER TABLE forecast_assignments ADD COLUMN created_at TEXT")
        con.execute("UPDATE forecast_assignments SET created_at=datetime('now') WHERE created_at IS NULL")
    cols = {r[1] for r in con.execute("PRAGMA table_info(forecast_assignments)")}
    if "session_start" not in cols:
        con.execute("ALTER TABLE forecast_assignments ADD COLUMN session_start TEXT")
        con.execute(
            "CREATE INDEX IF NOT EXISTS ix_forecast_session "
            "ON forecast_assignments(project_code, session_start)")
    _resplit_copied_day_hours(con)
    ensure_public_holidays(con)
    con.commit()
    con.close()


def _fold(name):
    if not name:
        return ""
    s = unicodedata.normalize("NFKD", str(name))
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return " ".join(s.upper().split())


def _norm_header(value):
    return re.sub(r"[^a-z0-9]+", " ", ("" if value is None else str(value)).strip().lower()).strip()


def _as_hours(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "")
    if not s:
        return None
    return float(s)


def _project_subtasks(con, project):
    rows = con.execute(
        """
        SELECT DISTINCT ts.id, ts.category_id, COALESCE(c.name, '') AS category, ts.name
          FROM task_sub_tasks ts
          LEFT JOIN categories c ON c.id = ts.category_id
         WHERE ts.category_id IN (
               SELECT DISTINCT ct.category_id FROM category_tasks ct
                 JOIN master_tasks mt ON mt.id = ct.master_task_id
                WHERE mt.project_code = ?)
         ORDER BY 3, ts.name
        """,
        (project,),
    ).fetchall()
    if rows:
        return rows
    return con.execute(
        """
        SELECT ts.id, ts.category_id, COALESCE(c.name, '') AS category, ts.name
          FROM task_sub_tasks ts
          LEFT JOIN categories c ON c.id = ts.category_id
         ORDER BY 3, ts.name
        """
    ).fetchall()


def _match_subtask(catalog, sub_name, category_name=""):
    want = _fold(sub_name)
    if not want:
        return None
    cat = _fold(category_name)
    hits = [r for r in catalog if _fold(r["name"]) == want]
    if cat:
        scoped = [r for r in hits if _fold(r["category"]) == cat]
        if scoped:
            hits = scoped
    if len(hits) == 1:
        return hits[0]
    return None


def _latest_status_map(con, project):
    """Latest daily-task status per sub-task, and per zone/level/sub-task."""
    rows = con.execute(
        """
        SELECT t.building_id, t.level_id, tl.task_sub_task_id,
               COALESCE(st.name, '') AS status, t.pct_complete
          FROM tasks t
          JOIN task_sub_task_links tl ON tl.task_id = t.id
          LEFT JOIN statuses st ON st.id = t.status_id
         WHERE t.project_code = ?
         ORDER BY t.work_date, t.id
        """,
        (project,),
    ).fetchall()
    by_sub, by_cell = {}, {}
    for r in rows:
        sid = r["task_sub_task_id"]
        by_sub[sid] = r["status"]
        by_cell[(r["building_id"], r["level_id"], sid)] = r["status"]
    return by_sub, by_cell


def _logged_hours_map(con, project):
    try:
        rows = con.execute(
            """
            SELECT tl.task_sub_task_id, ROUND(SUM(v.spent_hours), 2) AS hrs
              FROM v_task_effort v
              JOIN task_sub_task_links tl ON tl.task_id = v.task_id
             WHERE v.project_code = ?
             GROUP BY tl.task_sub_task_id
            """,
            (project,),
        ).fetchall()
    except sqlite3.Error:
        rows = con.execute(
            """
            SELECT tl.task_sub_task_id, ROUND(SUM(t.hours), 2) AS hrs
              FROM tasks t
              JOIN task_sub_task_links tl ON tl.task_id = t.id
             WHERE t.project_code = ?
             GROUP BY tl.task_sub_task_id
            """,
            (project,),
        ).fetchall()
    return {r["task_sub_task_id"]: float(r["hrs"] or 0) for r in rows}


def _status_bucket(status):
    s = (status or "").strip().lower()
    if s == "completed":
        return "completed"
    if s == "ongoing":
        return "ongoing"
    if s in ("on hold", "hold"):
        return "hold"
    if s in ("not started", "n", "todo"):
        return "not_started"
    if s:
        return "ongoing"
    return "not_started"


def project_progress(db_path, project):
    """Budget hours vs latest Completed / Ongoing status from daily tasks."""
    empty = {"rows": [], "total": 0.0, "completed": 0.0, "ongoing": 0.0,
             "hold": 0.0, "not_started": 0.0, "logged": 0.0}
    if not project:
        return empty
    con = connect(db_path)
    budget = con.execute(
        """
        SELECT b.id, b.hours, b.task_sub_task_id, b.building_id, b.level_id,
               COALESCE(c.name, '') AS category, ts.name AS sub_task,
               COALESCE(z.name, '') AS zone, COALESCE(l.label, '') AS level
          FROM forecast_subtask_budget b
          JOIN task_sub_tasks ts ON ts.id = b.task_sub_task_id
          LEFT JOIN categories c ON c.id = b.category_id
          LEFT JOIN buildings z ON z.id = b.building_id
          LEFT JOIN levels l ON l.id = b.level_id
         WHERE b.project_code = ?
         ORDER BY 6, ts.name, z.name, (l.number IS NULL), l.number
        """,
        (project,),
    ).fetchall()
    if not budget:
        con.close()
        return empty
    by_sub, by_cell = _latest_status_map(con, project)
    logged = _logged_hours_map(con, project)
    con.close()
    out = []
    totals = {"completed": 0.0, "ongoing": 0.0, "hold": 0.0, "not_started": 0.0, "logged": 0.0}
    logged_seen = set()
    for r in budget:
        hrs = float(r["hours"] or 0)
        sid = r["task_sub_task_id"]
        if r["building_id"] or r["level_id"]:
            status = by_cell.get((r["building_id"], r["level_id"], sid), "")
        else:
            status = by_sub.get(sid, "")
        bucket = _status_bucket(status)
        log_h = logged.get(sid, 0.0)
        totals[bucket] = totals.get(bucket, 0.0) + hrs
        if sid not in logged_seen:
            totals["logged"] += log_h
            logged_seen.add(sid)
        out.append({
            "category": r["category"] or "",
            "sub_task": r["sub_task"] or "",
            "zone": r["zone"] or "",
            "level": r["level"] or "",
            "hours": hrs,
            "status": status or "Not started",
            "bucket": bucket,
            "logged": log_h,
        })
    total = sum(float(r["hours"] or 0) for r in budget)
    return {
        "rows": out,
        "total": total,
        "completed": totals["completed"],
        "ongoing": totals["ongoing"],
        "hold": totals["hold"],
        "not_started": totals["not_started"],
        "logged": totals["logged"],
    }


def subtask_hours_workbook_bytes(db_path, project):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    con = connect(db_path)
    catalog = _project_subtasks(con, project)
    saved = {(r["task_sub_task_id"], r["building_id"] or 0, r["level_id"] or 0): r["hours"]
             for r in con.execute(
                 "SELECT task_sub_task_id, building_id, level_id, hours "
                 "FROM forecast_subtask_budget WHERE project_code=?",
                 (project,)).fetchall()}
    extra = con.execute(
        """
        SELECT b.hours, COALESCE(c.name,'') AS category, ts.name AS sub_task,
               COALESCE(z.name,'') AS zone, COALESCE(l.label,'') AS level
          FROM forecast_subtask_budget b
          JOIN task_sub_tasks ts ON ts.id = b.task_sub_task_id
          LEFT JOIN categories c ON c.id = b.category_id
          LEFT JOIN buildings z ON z.id = b.building_id
          LEFT JOIN levels l ON l.id = b.level_id
         WHERE b.project_code=? AND (b.building_id IS NOT NULL OR b.level_id IS NOT NULL)
         ORDER BY ts.name, z.name
        """,
        (project,),
    ).fetchall()
    con.close()

    wb = Workbook()
    notes = wb.active
    notes.title = "Instructions"
    notes["A1"] = "Hours to complete each sub-task"
    notes["A1"].font = Font(bold=True, size=14)
    lines = [
        "",
        "This sheet is the budget for project % completed and % ongoing.",
        "",
        "1. Open the Hours sheet. Every project sub-task is listed.",
        "2. Type Hours to complete for each sub-task (hours for that sub-task on this project).",
        "3. Leave Hours blank to skip a sub-task — it will not count in the %.",
        "4. Optional: fill Zone and Level if the hours are for one location only.",
        "5. Upload the file on Forecast. The new file replaces the last hours list for this project.",
        "6. Daily Task Entry status (Completed / Ongoing) is then weighed against these hours.",
        "   Completed hours / total hours = % completed. Ongoing hours / total hours = % ongoing.",
    ]
    for i, line in enumerate(lines, start=2):
        notes[f"A{i}"] = line
    notes.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Hours", 0)
    headers = ["Category", "Sub Task", "Time (hours)", "Zone", "Level"]
    fill = PatternFill("solid", fgColor=INK)
    hdr = Font(bold=True, color="FFFFFF")
    time_fill = PatternFill("solid", fgColor="FFF3CD")
    cat_fill = PatternFill("solid", fgColor="E8EEF3")
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.font = hdr
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 12
    ws.freeze_panes = "A2"
    r = 2
    last_cat = object()
    for row in catalog:
        hrs = saved.get((row["id"], 0, 0))
        cat = row["category"] or ""
        c_cat = ws.cell(r, 1, cat)
        ws.cell(r, 2, row["name"])
        c_hrs = ws.cell(r, 3, float(hrs) if hrs is not None else None)
        c_hrs.fill = time_fill
        c_hrs.number_format = "0.00"
        if cat != last_cat:
            c_cat.font = Font(bold=True)
            last_cat = cat
        else:
            c_cat.fill = cat_fill
        r += 1
    for row in extra:
        ws.cell(r, 1, row["category"] or "")
        ws.cell(r, 2, row["sub_task"])
        c_hrs = ws.cell(r, 3, float(row["hours"] or 0))
        c_hrs.fill = time_fill
        c_hrs.number_format = "0.00"
        ws.cell(r, 4, row["zone"] or "")
        ws.cell(r, 5, row["level"] or "")
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return f"subtask_hours_{project}.xlsx", buf.getvalue()


def import_subtask_hours(db_path, xlsx_path, project):
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    sheet = wb["Hours"] if "Hours" in wb.sheetnames else wb.active
    rows = list(sheet.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("The workbook has no rows.")
    headers = [_norm_header(v) for v in rows[0]]

    def col(*names):
        for n in names:
            if n in headers:
                return headers.index(n)
        return None

    c_cat = col("category")
    c_sub = col("sub task", "subtask", "sub-task")
    c_hrs = col("hours to complete", "hours", "hour", "estimated hours", "hrs",
                "time", "time hours", "time (hours)")
    c_zone = col("zone", "building")
    c_level = col("level", "floor")
    if c_sub is None or c_hrs is None:
        raise ValueError("Need columns Sub Task and Hours to complete.")

    con = connect(db_path)
    catalog = _project_subtasks(con, project)
    zones = {_fold(r["name"]): r["id"] for r in con.execute(
        "SELECT id, name FROM buildings WHERE project_code=?", (project,))}
    levels = {_fold(r["label"]): r["id"] for r in con.execute(
        "SELECT l.id, l.label FROM levels l JOIN project_levels pl ON pl.level_id=l.id "
        "WHERE pl.project_code=?", (project,))}

    saved, errors = [], []
    seen = set()
    for i, row in enumerate(rows[1:], start=2):
        def val(idx):
            if idx is None or idx >= len(row):
                return ""
            v = row[idx]
            return "" if v is None else str(v).strip()

        sub = val(c_sub)
        if not sub or _fold(sub) in ("SUB TASK",):
            continue
        try:
            hrs = _as_hours(row[c_hrs] if c_hrs < len(row) else None)
        except (TypeError, ValueError):
            errors.append(f"Row {i}: hours is not a number ({sub}).")
            continue
        if hrs is None:
            continue
        if hrs < 0:
            errors.append(f"Row {i}: hours cannot be negative ({sub}).")
            continue
        hit = _match_subtask(catalog, sub, val(c_cat))
        if not hit:
            errors.append(f"Row {i}: sub-task not found ({sub}).")
            continue
        zone_id = levels_id = None
        zone_name, level_name = val(c_zone), val(c_level)
        if zone_name:
            zone_id = zones.get(_fold(zone_name))
            if not zone_id:
                errors.append(f"Row {i}: zone not found ({zone_name}).")
                continue
        if level_name:
            levels_id = levels.get(_fold(level_name))
            if not levels_id:
                errors.append(f"Row {i}: level not found ({level_name}).")
                continue
        key = (hit["id"], zone_id or 0, levels_id or 0)
        if key in seen:
            errors.append(f"Row {i}: duplicate {sub}.")
            continue
        seen.add(key)
        saved.append((project, hit["category_id"], hit["id"], zone_id, levels_id, hrs))

    con.execute("DELETE FROM forecast_subtask_budget WHERE project_code=?", (project,))
    con.executemany(
        "INSERT INTO forecast_subtask_budget"
        "(project_code, category_id, task_sub_task_id, building_id, level_id, hours) "
        "VALUES (?,?,?,?,?,?)",
        saved,
    )
    con.commit()
    con.close()
    return {"saved": len(saved), "errors": errors}


def _pct(part, total):
    if total <= 0:
        return 0.0
    return round(100.0 * part / total, 1)


def _progress_html(prog, qs, project, can_write):
    total = prog.get("total") or 0.0
    rows = prog.get("rows") or []
    pc = _pct(prog["completed"], total)
    po = _pct(prog["ongoing"], total)
    ph = _pct(prog["hold"], total)
    pn = _pct(prog["not_started"], total)
    upload = ""
    if can_write:
        upload = (
            f'<form method="POST" action="/forecast/subtask-hours" enctype="multipart/form-data" '
            f'class="actions" style="margin-top:12px">'
            f'<input type="hidden" name="qs" value="{html.escape(qs)}">'
            f'<input type="hidden" name="project" value="{html.escape(project)}">'
            f'<input type="file" name="xlsx" accept=".xlsx" required>'
            f'<button type="submit">Upload hours Excel</button></form>'
        )
    if not rows:
        return (
            f'<div class="card">'
            f'<h2 style="margin-top:0">Project % — hours to complete</h2>'
            f'<p class="hint" style="margin-top:0">Upload an Excel of each sub-task and how many hours '
            f'it takes to finish. That list is the basis for % completed and % ongoing '
            f'(from Daily Task Entry status).</p>'
            f'<div class="exports">'
            f'<a class="btn grey" href="/forecast/subtask-hours.xlsx?project={html.escape(project)}">'
            f'Download hours template</a></div>'
            f'{upload}</div>'
        )
    bits = []
    if pc:
        bits.append(f'<span class="done" style="flex:{pc}">{pc:g}% completed</span>')
    if po:
        bits.append(f'<span class="wip" style="flex:{po}">{po:g}% ongoing</span>')
    rest = pn + ph
    if rest:
        label = f"{pn:g}% not started" + (f" · {ph:g}% on hold" if ph else "")
        bits.append(f'<span class="todo" style="flex:{rest}">{html.escape(label)}</span>')
    if not bits:
        bits.append('<span class="todo" style="flex:100">No hours yet</span>')
    bar = f'<div class="prog">{"".join(bits)}</div>'
    stats = (
        f'<div class="progstats">'
        f'<div>Completed <b>{pc:g}%</b><div class="hint">{_h(prog["completed"])} of {_h(total)} hrs</div></div>'
        f'<div>Ongoing <b>{po:g}%</b><div class="hint">{_h(prog["ongoing"])} hrs</div></div>'
        f'<div>Not started <b>{pn:g}%</b><div class="hint">{_h(prog["not_started"])} hrs</div></div>'
        f'<div>Logged so far <b>{_h(prog["logged"])}</b><div class="hint">hours on daily tasks</div></div>'
        f'</div>'
    )
    body = []
    for r in rows:
        cls = {"completed": "okh", "ongoing": "short", "hold": "short"}.get(r["bucket"], "")
        loc = " · ".join(x for x in (r.get("zone"), r.get("level")) if x)
        body.append(
            "<tr>"
            f'<td>{html.escape(r["category"])}</td>'
            f'<td>{html.escape(r["sub_task"])}'
            f'{(" <span class=hint>(" + html.escape(loc) + ")</span>") if loc else ""}</td>'
            f'<td class="num">{_h(r["hours"])}</td>'
            f'<td class="{cls}">{html.escape(r["status"])}</td>'
            f'<td class="num">{_h(r["logged"])}</td></tr>'
        )
    table = (
        '<div class="scroll"><table><thead><tr>'
        "<th>Category</th><th>Sub Task</th><th class='num'>Hours to complete</th>"
        "<th>Latest status</th><th class='num'>Logged hrs</th>"
        f"</tr></thead><tbody>{''.join(body)}</tbody></table></div>"
    )
    return (
        f'<div class="card">'
        f'<h2 style="margin-top:0">Project % — hours to complete</h2>'
        f'<p class="hint" style="margin-top:0">Basis is the uploaded hours list. Status comes from '
        f'the latest Daily Task Entry for each sub-task (Completed / Ongoing / not started).</p>'
        f'{bar}{stats}'
        f'<div class="exports">'
        f'<a class="btn grey" href="/forecast/subtask-hours.xlsx?project={html.escape(project)}">'
        f'Download hours Excel</a>'
        f'<a class="btn green" href="/forecast/export?{qs}&kind=progress">Export progress</a>'
        f'</div>'
        f'{upload}{table}</div>'
    )


def _resplit_copied_day_hours(con):
    """Old saves copied a full day onto every ticked sub-task (e.g. 6 x 8.25 = 49.5).
    Those rows share one working day instead."""
    groups = con.execute(
        """
        SELECT person_id, work_date
          FROM forecast_assignments
         GROUP BY person_id, work_date
        HAVING SUM(estimated_hours) > ? + 0.001
        """,
        (HOURS_PER_DAY,),
    ).fetchall()
    for g in groups:
        rows = con.execute(
            """
            SELECT id, estimated_hours FROM forecast_assignments
             WHERE person_id=? AND work_date=? ORDER BY id
            """,
            (g["person_id"], g["work_date"]),
        ).fetchall()
        if not rows:
            continue
        total = sum(float(r["estimated_hours"] or 0) for r in rows)
        if total <= 0:
            continue
        assigned = 0.0
        last = len(rows) - 1
        for i, r in enumerate(rows):
            if i == last:
                new_h = round(HOURS_PER_DAY - assigned, 2)
            else:
                new_h = round(float(r["estimated_hours"] or 0) / total * HOURS_PER_DAY, 2)
                assigned += new_h
            if new_h < 0.01:
                new_h = 0.01
            con.execute(
                "UPDATE forecast_assignments SET estimated_hours=? WHERE id=?",
                (new_h, r["id"]),
            )


def _one(form, name):
    v = (form.get(name, [""])[0] or "").strip()
    return v or None


def _int(form, name):
    v = _one(form, name)
    if v is None:
        return None
    try:
        return int(v)
    except ValueError:
        return None


def _float(form, name):
    v = _one(form, name)
    if v is None:
        return None
    try:
        return float(v)
    except ValueError:
        return None


def _ids(form, name):
    out = []
    for v in form.get(name, []):
        v = (v or "").strip()
        if not v:
            continue
        try:
            out.append(int(v))
        except ValueError:
            continue
    return out


def _dates(form):
    ds = [(v or "").strip() for v in form.get("work_dates", [])]
    ds = [d for d in ds if d]
    if not ds:
        one = _one(form, "work_date")
        if one:
            ds = [one]
    for d in ds:
        dt.date.fromisoformat(d)
    return ds


def _axis(ids, required, label):
    if required and not ids:
        raise ValueError(f"Tick at least one {label}.")
    return ids or [None]


def _work_items(form, db_path):
    """Each ticked sub-task is one work item. Otherwise category × task, else one blank."""
    subs = _ids(form, "task_sub_task_ids")
    cats = _ids(form, "category_ids")
    tasks = _ids(form, "task_name_ids")
    if subs:
        con = connect(db_path)
        rows = {r["id"]: r for r in con.execute(
            "SELECT id, category_id, task_name_id FROM task_sub_tasks").fetchall()}
        con.close()
        items = []
        for sid in subs:
            r = rows.get(sid)
            if r:
                items.append((r["category_id"], r["task_name_id"], sid))
        return items or [(None, None, None)]
    if cats and tasks:
        return [(c, t, None) for c in cats for t in tasks]
    if cats:
        t = tasks[0] if tasks else None
        return [(c, t, None) for c in cats]
    if tasks:
        c = cats[0] if cats else None
        return [(c, t, None) for t in tasks]
    return [(None, None, None)]


def window_from(q, ref, db_path=None):
    codes = [p["code"] for p in ref.get("projects") or []]
    project = (q.get("project", [""])[0] or "").strip()
    if project not in codes:
        project = codes[0] if codes else ""
    start = (q.get("start", [""])[0] or "").strip()
    start_given = bool(start)
    try:
        dt.date.fromisoformat(start)
    except ValueError:
        start = dt.date.today().isoformat()
        start_given = False
    try:
        days = int(q.get("days", ["5"])[0])
    except (TypeError, ValueError):
        days = 5
    days = max(1, min(92, days))
    end = (q.get("end", [""])[0] or "").strip()
    if end:
        try:
            e = dt.date.fromisoformat(end)
            s = dt.date.fromisoformat(start)
            if e < s:
                s, e = e, s
                start = s.isoformat()
            days = min(92, max(1, (e - s).days + 1))
        except ValueError:
            pass
    if not start_given and not end and db_path and project:
        start, days = default_window(db_path, project, days)
    return project, start, days


def resolve_person_id(db_path, token):
    """Emp ID, person id, or unique name. None = no filter."""
    token = (token or "").strip()
    if not token:
        return None
    con = connect(db_path)
    if token.isdigit():
        row = con.execute("SELECT id FROM people WHERE id=?", (int(token),)).fetchone()
        if row:
            con.close()
            return row["id"]
        row = con.execute(
            "SELECT id FROM people WHERE emp_code=? COLLATE NOCASE", (token,)).fetchone()
        if row:
            con.close()
            return row["id"]
    row = con.execute(
        "SELECT id FROM people WHERE emp_code=? COLLATE NOCASE", (token,)).fetchone()
    if row:
        con.close()
        return row["id"]
    row = con.execute(
        "SELECT id FROM people WHERE name=? COLLATE NOCASE", (token,)).fetchone()
    if row:
        con.close()
        return row["id"]
    want = _compact_id(token)
    hits = [r for r in con.execute("SELECT id, name, emp_code FROM people")
            if _compact_id(r["name"]) == want or _compact_id(r["emp_code"] or "") == want]
    con.close()
    return hits[0]["id"] if len(hits) == 1 else None


def person_from(q):
    return (q.get("person", [""])[0] or "").strip()


def date_list(start, days):
    d0 = dt.date.fromisoformat(start)
    return [(d0 + dt.timedelta(days=i)).isoformat() for i in range(days)]


def window_end(start, days):
    return date_list(start, days)[-1]


def session_scope_sql(alias="f"):
    """This planning session, or any row whose work date sits in the window.

    session_start matches even if floor dates run past the visible To date.
    Work dates in the From–To range match even if they were stamped on an
    earlier session (so Export this plan is not an empty sheet).
    """
    return (
        f"({alias}.session_start = ? OR "
        f"({alias}.work_date >= ? AND {alias}.work_date <= ?))"
    )


def plan_window_or_latest(db_path, project, start, days):
    """Keep the chosen window when it has rows; otherwise the latest saved plan."""
    end = window_end(start, days)
    if not project:
        return start, days, end
    con = connect(db_path)
    n = con.execute(
        "SELECT COUNT(*) FROM forecast_assignments f "
        "WHERE f.project_code=? AND " + session_scope_sql("f"),
        (project, start, start, end),
    ).fetchone()[0]
    if n:
        con.close()
        return start, days, end
    row = con.execute(
        """
        SELECT session_start AS s, MIN(work_date) AS mn, MAX(work_date) AS mx
          FROM forecast_assignments
         WHERE project_code=? AND session_start IS NOT NULL AND TRIM(session_start)<>''
         GROUP BY session_start
         ORDER BY MAX(work_date) DESC
         LIMIT 1
        """,
        (project,),
    ).fetchone()
    if not row:
        row = con.execute(
            """
            SELECT MIN(work_date) AS mn, MAX(work_date) AS mx
              FROM forecast_assignments WHERE project_code=?
            """,
            (project,),
        ).fetchone()
        con.close()
        if not row or not row["mn"]:
            return start, days, end
        s, e = row["mn"], row["mx"]
    else:
        con.close()
        s, e = row["s"] or row["mn"], row["mx"]
    try:
        ndays = max(1, (dt.date.fromisoformat(e) - dt.date.fromisoformat(s)).days + 1)
    except (TypeError, ValueError):
        return start, days, end
    return s, ndays, e


def window_qs(project, start, days, extra="", person="", end=""):
    from urllib.parse import quote
    q = f"project={quote(project or '')}&start={quote(start or '')}&days={days}"
    if end:
        q += f"&end={quote(end)}"
    if person:
        q += f"&person={quote(person)}"
    return q + (("&" + extra) if extra else "")


def saved_periods(db_path, project):
    """One chip per loading-plan session (or a contiguous date run for old rows)."""
    if not project:
        return []
    con = connect(db_path)
    sessions = con.execute(
        """
        SELECT session_start,
               MIN(work_date) AS first_day,
               MAX(work_date) AS last_day,
               COUNT(*) AS rows
          FROM forecast_assignments
         WHERE project_code=? AND session_start IS NOT NULL
           AND TRIM(session_start)<>''
         GROUP BY session_start
         ORDER BY MIN(work_date)
        """,
        (project,),
    ).fetchall()
    null_dates = [r[0] for r in con.execute(
        "SELECT DISTINCT work_date FROM forecast_assignments "
        "WHERE project_code=? AND (session_start IS NULL OR TRIM(session_start)='') "
        "ORDER BY work_date",
        (project,),
    ).fetchall()]
    null_counts = {r[0]: r[1] for r in con.execute(
        "SELECT work_date, COUNT(*) FROM forecast_assignments "
        "WHERE project_code=? AND (session_start IS NULL OR TRIM(session_start)='') "
        "GROUP BY work_date",
        (project,),
    ).fetchall()}
    con.close()
    periods = []
    for r in sessions:
        first, last = r["first_day"], r["last_day"]
        try:
            ndays = max(1, (dt.date.fromisoformat(last) - dt.date.fromisoformat(first)).days + 1)
        except ValueError:
            ndays = 1
        periods.append({
            "start": r["session_start"] or first,
            "end": last,
            "days": ndays,
            "rows": r["rows"],
            "session_start": r["session_start"],
        })
    if null_dates:
        run_start = run_end = null_dates[0]
        n = null_counts.get(null_dates[0], 0)
        for iso in null_dates[1:]:
            prev = dt.date.fromisoformat(run_end)
            cur = dt.date.fromisoformat(iso)
            if cur == prev + dt.timedelta(days=1):
                run_end = iso
                n += null_counts.get(iso, 0)
                continue
            ndays = (dt.date.fromisoformat(run_end) - dt.date.fromisoformat(run_start)).days + 1
            periods.append({"start": run_start, "end": run_end, "days": ndays, "rows": n,
                            "session_start": ""})
            run_start = run_end = iso
            n = null_counts.get(iso, 0)
        ndays = (dt.date.fromisoformat(run_end) - dt.date.fromisoformat(run_start)).days + 1
        periods.append({"start": run_start, "end": run_end, "days": ndays, "rows": n,
                        "session_start": ""})
    periods.sort(key=lambda p: p["start"] or "")
    return periods


def saved_plan_fields(db_path, project, start, days):
    """Rebuild Add-floors ticks from the assignments in this week’s plan."""
    if not project or not start:
        return None
    end = window_end(start, days)
    con = connect(db_path)
    rows = con.execute(
        """
        SELECT person_id, work_date, level_id, category_id, building_id, model_id,
               drawing_number, session_start
          FROM forecast_assignments f
         WHERE f.project_code=? AND """ + session_scope_sql("f") + """
         ORDER BY f.work_date, f.id
        """,
        (project, start, start, end),
    ).fetchall()
    if not rows:
        con.close()
        return None
    dates_by_floor = {}
    people, cats = set(), set()
    work_dates = []
    extra_names = []
    session = None
    model_id = zone_id = None
    for r in rows:
        session = session or r["session_start"]
        wd = r["work_date"]
        if wd and wd not in work_dates:
            work_dates.append(wd)
        if r["level_id"]:
            dates_by_floor.setdefault(int(r["level_id"]), set()).add(wd)
        if r["person_id"]:
            people.add(int(r["person_id"]))
        if r["category_id"]:
            cats.add(int(r["category_id"]))
        if model_id is None and r["model_id"]:
            model_id = int(r["model_id"])
        if zone_id is None and r["building_id"]:
            zone_id = int(r["building_id"])
        extra = (r["drawing_number"] or "").strip()
        if extra:
            extra_names.extend(x.strip() for x in extra.split(",") if x.strip())
    if extra_names:
        names = list(dict.fromkeys(extra_names))
        marks = ",".join("?" * len(names))
        for r in con.execute(
            f"SELECT id FROM categories WHERE name IN ({marks})", names
        ):
            cats.add(int(r["id"]))
    con.close()
    return {
        "start": work_dates[0] if work_dates else start,
        "session_start": session or start,
        "floor_days": {lid: len(ds) for lid, ds in dates_by_floor.items()},
        "people": people,
        "category_ids": cats,
        "model_id": model_id,
        "building_id": zone_id,
    }


def floor_days_from_q(q):
    out = {}
    for key, values in (q or {}).items():
        if not str(key).startswith("floor_days_"):
            continue
        try:
            lid = int(str(key)[11:])
        except ValueError:
            continue
        raw = (values[0] if values else "").strip()
        if not raw:
            continue
        try:
            n = int(raw)
        except ValueError:
            continue
        if n > 0:
            out[lid] = min(30, n)
    return out


def floor_days_extra(form):
    from urllib.parse import quote
    bits = []
    for key, values in form.items():
        if not str(key).startswith("floor_days_"):
            continue
        raw = (values[0] if values else "").strip()
        if raw:
            bits.append(quote(str(key), safe="") + "=" + quote(raw, safe=""))
    fill = (form.get("fill_floor_days", [""])[0] or "").strip()
    if fill:
        bits.append("fill_floor_days=" + quote(fill))
    plan_start = (form.get("plan_start", [""])[0] or "").strip()
    if plan_start:
        bits.append("plan_start=" + quote(plan_start))
    skip = "1" if "1" in [(v or "").strip() for v in form.get("skip_weekend", [])] else "0"
    if form.get("skip_weekend") is not None:
        bits.append("skip_weekend=" + skip)
    return "&".join(bits)


def delete_loading_plan(db_path, project, session_start=None, start=None, end=None):
    """Remove a saved loading plan. Logged daily tasks are not touched."""
    if not project:
        raise ValueError("Choose a project.")
    con = connect(db_path)
    if session_start:
        cur = con.execute(
            "DELETE FROM forecast_assignments WHERE project_code=? AND session_start=?",
            (project, session_start))
    elif start and end:
        cur = con.execute(
            "DELETE FROM forecast_assignments WHERE project_code=? AND work_date>=? AND work_date<=?"
            " AND (session_start IS NULL OR TRIM(session_start)='')",
            (project, start, end))
    else:
        con.close()
        raise ValueError("Choose a plan to remove.")
    n = cur.rowcount
    con.commit()
    con.close()
    return n


def default_window(db_path, project, days_fallback=5):
    """Open the plan that covers today, otherwise the latest saved plan.

    Jumping past the last assignment to empty future days hid the forecast
    list (and made Export loading sheet look broken). Use Plan next N days
    when you want a new window.
    """
    today = dt.date.today()
    periods = saved_periods(db_path, project)
    for p in periods:
        s, e = dt.date.fromisoformat(p["start"]), dt.date.fromisoformat(p["end"])
        if s <= today <= e:
            return p["start"], p["days"]
    if periods:
        last = periods[-1]
        return last["start"], last["days"]
    return today.isoformat(), days_fallback


def next_plan_start(db_path, project, start, days):
    periods = saved_periods(db_path, project)
    if periods:
        last = dt.date.fromisoformat(periods[-1]["end"])
        return (last + dt.timedelta(days=1)).isoformat()
    return (dt.date.fromisoformat(window_end(start, days)) + dt.timedelta(days=1)).isoformat()


def _actual_hours_by_person_day(con, project, start, end):
    from export_reports import not_manager_sql
    skip_mgr = not_manager_sql()
    skip_mgr_t = not_manager_sql("ta.person_id", "t.project_code")
    try:
        rows = con.execute(
            f"""
            SELECT person_id, work_date, ROUND(SUM(spent_hours), 2) AS hrs
              FROM v_task_effort
             WHERE project_code=? AND work_date>=? AND work_date<=?
               AND {skip_mgr}
             GROUP BY person_id, work_date
            """,
            (project, start, end),
        ).fetchall()
    except sqlite3.Error:
        rows = con.execute(
            f"""
            SELECT ta.person_id, t.work_date, ROUND(SUM(t.hours), 2) AS hrs
              FROM tasks t
              JOIN task_assignees ta ON ta.task_id=t.id
             WHERE t.project_code=? AND t.work_date>=? AND t.work_date<=?
               AND {skip_mgr_t}
             GROUP BY ta.person_id, t.work_date
            """,
            (project, start, end),
        ).fetchall()
    return {(r["person_id"], r["work_date"]): float(r["hrs"] or 0) for r in rows}


def _logged_work(con, project, start, end):
    """(person_id, work_date, sub_task_id) logged on daily tasks."""
    rows = con.execute(
        """
        SELECT ta.person_id, t.work_date, tl.task_sub_task_id
          FROM tasks t
          JOIN task_assignees ta ON ta.task_id=t.id
          JOIN task_sub_task_links tl ON tl.task_id=t.id
         WHERE t.project_code=? AND t.work_date>=? AND t.work_date<=?
        """,
        (project, start, end),
    ).fetchall()
    return {(r["person_id"], r["work_date"], r["task_sub_task_id"]) for r in rows}


def compare_forecast_actual(db_path, project, start, days, person_id=None):
    """Forecast hours vs daily-task hours for the same people and dates. Reads tasks, never writes."""
    if not project:
        return {"people": [], "items": [], "start": start, "end": ""}
    end = window_end(start, days)
    con = connect(db_path)
    who_sql, who_params = "", [project, start, end]
    if person_id:
        who_sql, who_params = " AND f.person_id=?", [project, start, end, person_id]
    fc = con.execute(
        """
        SELECT f.person_id, pe.name AS person, pe.emp_code, f.work_date,
               ROUND(SUM(f.estimated_hours), 2) AS hrs
          FROM forecast_assignments f
          JOIN people pe ON pe.id=f.person_id
         WHERE f.project_code=? AND f.work_date>=? AND f.work_date<=?""" + who_sql + """
         GROUP BY f.person_id, f.work_date
        """,
        who_params,
    ).fetchall()
    actual = _actual_hours_by_person_day(con, project, start, end)
    if person_id:
        actual = {k: v for k, v in actual.items() if k[0] == person_id}
    logged = _logged_work(con, project, start, end)
    items = con.execute(
        """
        SELECT f.work_date, pe.name AS person, pe.emp_code, c.name AS category,
               tn.name AS task, ts.name AS sub_task, f.person_id, f.task_sub_task_id,
               ROUND(f.estimated_hours, 2) AS hrs
          FROM forecast_assignments f
          JOIN people pe ON pe.id=f.person_id
          LEFT JOIN categories c ON c.id=f.category_id
          LEFT JOIN task_names tn ON tn.id=f.task_name_id
          LEFT JOIN task_sub_tasks ts ON ts.id=f.task_sub_task_id
         WHERE f.project_code=? AND f.work_date>=? AND f.work_date<=?""" + who_sql + """
         ORDER BY f.work_date, pe.name, f.id
        """,
        who_params,
    ).fetchall()

    by_person = {}
    people_days = []
    for r in fc:
        pid, d, fhrs = r["person_id"], r["work_date"], float(r["hrs"] or 0)
        ahrs = actual.get((pid, d), 0.0)
        people_days.append({
            "person": r["person"], "emp_code": r["emp_code"] or "", "work_date": d,
            "forecast": fhrs, "actual": ahrs, "diff": round(ahrs - fhrs, 2),
        })
        rec = by_person.setdefault(pid, {
            "person": r["person"], "emp_code": r["emp_code"] or "",
            "forecast": 0.0, "actual": 0.0,
        })
        rec["forecast"] = round(rec["forecast"] + fhrs, 2)
    seen = set(by_person)
    for (pid, d), ahrs in actual.items():
        if pid in by_person:
            by_person[pid]["actual"] = round(by_person[pid]["actual"] + ahrs, 2)
        else:
            who = con.execute("SELECT name, emp_code FROM people WHERE id=?", (pid,)).fetchone()
            if not who:
                continue
            by_person[pid] = {
                "person": who["name"], "emp_code": who["emp_code"] or "",
                "forecast": 0.0, "actual": ahrs,
            }
            seen.add(pid)
    people = []
    for rec in sorted(by_person.values(), key=lambda x: x["person"]):
        rec = dict(rec)
        rec["diff"] = round(rec["actual"] - rec["forecast"], 2)
        people.append(rec)

    item_rows = []
    for r in items:
        done = (r["person_id"], r["work_date"], r["task_sub_task_id"]) in logged if r["task_sub_task_id"] else False
        item_rows.append({
            "work_date": r["work_date"], "person": r["person"], "emp_code": r["emp_code"] or "",
            "category": r["category"] or "", "task": r["task"] or "",
            "sub_task": r["sub_task"] or "", "forecast": float(r["hrs"] or 0),
            "logged": done,
        })
    extra_sql = ""
    extra_params = [project, start, end]
    if person_id:
        extra_sql = " AND ta.person_id=?"
        extra_params.append(person_id)
    extra_logged = con.execute(
        """
        SELECT t.work_date, pe.name AS person, ts.name AS sub_task, tn.name AS task
          FROM tasks t
          JOIN task_assignees ta ON ta.task_id=t.id
          JOIN people pe ON pe.id=ta.person_id
          JOIN task_sub_task_links tl ON tl.task_id=t.id
          JOIN task_sub_tasks ts ON ts.id=tl.task_sub_task_id
          LEFT JOIN task_names tn ON tn.id=t.task_name_id
         WHERE t.project_code=? AND t.work_date>=? AND t.work_date<=?""" + extra_sql + """
           AND NOT EXISTS (
               SELECT 1 FROM forecast_assignments f
                WHERE f.project_code=t.project_code AND f.person_id=ta.person_id
                  AND f.work_date=t.work_date AND f.task_sub_task_id=tl.task_sub_task_id)
         ORDER BY t.work_date, pe.name
        """,
        extra_params,
    ).fetchall()
    con.close()
    for r in extra_logged:
        item_rows.append({
            "work_date": r["work_date"], "person": r["person"], "emp_code": "",
            "category": "", "task": r["task"] or "",
            "sub_task": r["sub_task"] or "", "forecast": 0.0, "logged": True, "extra": True,
        })
    return {"people": people, "days": people_days, "items": item_rows, "start": start, "end": end}


def project_row(ref, code):
    return next((p for p in ref.get("projects") or [] if p["code"] == code), None)


def _has_left(row):
    return bool(str(row.get("left_on") or "").strip())


def project_team_configured(ref, project):
    return any(x.get("project_code") == project for x in (ref.get("project_people") or []))


def without_former_members(ref):
    """Drop people who already left a project, so Forecast pickers cannot see them."""
    out = dict(ref)
    today = dt.date.today().isoformat()
    current = []
    active = set()
    for x in ref.get("project_people") or []:
        if _has_left(x):
            continue
        joined = str(x.get("joined_on") or "").strip()
        if joined and joined > today:
            continue
        current.append(x)
        active.add((x.get("project_code"), str(x.get("person_id"))))
    out["project_people"] = current
    out["lead_support"] = [
        r for r in (ref.get("lead_support") or [])
        if (r.get("project_code"), str(r.get("lead_id"))) in active
        and (r.get("project_code"), str(r.get("support_id"))) in active
    ]
    return out


def current_team_ids(ref, project, on_date=None):
    """Person IDs currently on the project team.

    People taken off the team (`left_on` set) and people whose join date is
    still in the future are omitted. Returns ([], False) when the project has
    never had a team — callers may then fall back to everyone.
    """
    on_date = on_date or dt.date.today().isoformat()
    ids = []
    seen = set()
    any_team = False
    for x in ref.get("project_people") or []:
        if x.get("project_code") != project:
            continue
        any_team = True
        if _has_left(x):
            continue
        joined = str(x.get("joined_on") or "").strip()
        if joined and joined > on_date:
            continue
        pid = str(x["person_id"])
        if pid in seen:
            continue
        seen.add(pid)
        ids.append(pid)
    return ids, any_team


def _active_team_ids_db(con, project, on_date=None):
    has = con.execute(
        "SELECT 1 FROM project_people WHERE project_code=? LIMIT 1",
        (project,)).fetchone()
    if not has:
        return None
    on_date = on_date or dt.date.today().isoformat()
    return {
        int(r["person_id"])
        for r in con.execute(
            "SELECT person_id FROM project_people "
            "WHERE project_code=? AND left_on IS NULL "
            "AND (joined_on IS NULL OR joined_on<=?)",
            (project, on_date),
        )
    }


def list_assignments(db_path, project, start, days, person_id=None):
    if not project:
        return []
    end = window_end(start, days)
    con = connect(db_path)
    sql = """
        SELECT f.id, f.work_date, f.project_code, f.service, f.person_id, f.category_id, f.building_id,
               f.level_id, f.area, f.estimated_hours, f.task_name_id, f.task_sub_task_id,
               f.model_id, f.drawing_name, f.drawing_number, f.scale, f.paper_size, f.notes,
               pe.name AS person, pe.emp_code,
               c.name AS category, b.name AS zone, l.label AS level,
               tn.name AS task, ts.name AS sub_task, m.name AS model
        FROM forecast_assignments f
        JOIN people pe ON pe.id = f.person_id
        LEFT JOIN categories c ON c.id = f.category_id
        LEFT JOIN buildings b ON b.id = f.building_id
        LEFT JOIN levels l ON l.id = f.level_id
        LEFT JOIN task_names tn ON tn.id = f.task_name_id
        LEFT JOIN task_sub_tasks ts ON ts.id = f.task_sub_task_id
        LEFT JOIN models m ON m.id = f.model_id
        WHERE f.project_code = ? AND """ + session_scope_sql("f")
    params = [project, start, start, end]
    if person_id:
        sql += " AND f.person_id=?"
        params.append(person_id)
    sql += " ORDER BY f.work_date, pe.name, f.id"
    rows = con.execute(sql, params).fetchall()
    con.close()
    return [dict(r) for r in rows]


def get_assignment(db_path, assignment_id):
    con = connect(db_path)
    r = con.execute("SELECT * FROM forecast_assignments WHERE id=?", (int(assignment_id),)).fetchone()
    con.close()
    return dict(r) if r else None


def save_assignment(db_path, form, entered_by):
    """One row per person per day on a single floor. Edit updates a single row."""
    ensure_forecast_schema(db_path)
    project = _one(form, "project")
    if not project:
        raise ValueError("Choose a project.")
    service = _one(form, "service")
    if service not in SERVICES:
        raise ValueError("Service must be Modelling Works or Sheet Work.")
    hours = _float(form, "estimated_hours")
    if hours is None or hours <= 0:
        raise ValueError("Hours must be greater than 0.")
    area = _float(form, "area")
    people = _ids(form, "person_ids") or ([_int(form, "person_id")] if _int(form, "person_id") else [])
    people = [p for p in people if p]
    if not people:
        raise ValueError("Tick at least one person.")
    dates = _dates(form)
    if not dates:
        raise ValueError("Tick at least one work date.")
    session_start = _one(form, "session_start") or _one(form, "start") or dates[0]
    try:
        dt.date.fromisoformat(session_start)
    except ValueError:
        session_start = dates[0]
    edit_id = _one(form, "edit_id")
    if edit_id:
        if hours > HOURS_PER_DAY + 0.001:
            raise ValueError(
                f"One row is one person on one day. Cap is {HOURS_PER_DAY:g} hrs "
                f"(09:00–18:30 minus 75 min breaks)."
            )
    else:
        person_days = len(people) * len(dates)
        max_total = round(HOURS_PER_DAY * person_days, 2)
        if hours > max_total + 0.001:
            raise ValueError(
                f"Too many hours. {len(people)} person(s) × {len(dates)} day(s) can be at most "
                f"{max_total:g} hrs. A working day is {HOURS_PER_DAY:g} hrs "
                f"(09:00–18:30 minus 75 min breaks). Sub-tasks share that day — they do not each get "
                f"{HOURS_PER_DAY:g} hrs."
            )

    con = connect(db_path)
    prow = con.execute("SELECT no_zone FROM projects WHERE code=?", (project,)).fetchone()
    allowed = _active_team_ids_db(con, project)
    con.close()
    if allowed is not None:
        people = [p for p in people if p in allowed]
        if not people:
            raise ValueError(
                "Those people are not on this project's team. "
                "Add them in Manage Lists → Project team.")
    use_zone = not (prow and prow["no_zone"])

    zone_ids = _ids(form, "building_ids")
    if use_zone and not zone_ids:
        raise ValueError("Tick at least one Zone.")
    zone_id = zone_ids[0] if zone_ids else None
    levels = _ids(form, "level_ids")
    if not levels:
        raise ValueError("Tick at least one Floor.")
    if len(levels) > 1:
        n = len(people) * len(levels)
        raise ValueError(
            f"Assign work is one floor at a time. You ticked {len(people)} people and "
            f"{len(levels)} floors. Use Create loading plan "
            f"({len(people)} × days on each floor — "
            f"{n} rows if each floor is 1 day). Models and sub-tasks are not multiplied.")
    level_id = levels[0]
    if service == "Modelling Works":
        model_ids = _ids(form, "model_ids")
        model_id = model_ids[0] if model_ids else None
        drawing_name = drawing_number = scale = paper_size = None
    else:
        model_id = None
        drawing_name = _one(form, "drawing_name")
        if not drawing_name:
            raise ValueError("Drawing name is required for Sheet Work.")
        drawing_number = _one(form, "drawing_number")
        scale = _one(form, "scale")
        paper_size = _one(form, "paper_size")

    work_items = _work_items(form, db_path)
    work = work_items[0] if work_items else (None, None, None)

    def row_cols(work_date, person_id, zone_id, level_id, model_id, work, row_hours,
                 model_names, sub_names):
        cat_id, task_id, sub_id = work
        mname = model_names.get(int(model_id)) if model_id else None
        sname = sub_names.get(int(sub_id)) if sub_id else None
        if service == "Modelling Works":
            dname, dnum = mname, sname
        else:
            dname, dnum = drawing_name, (drawing_number or sname)
        return {
            "work_date": work_date,
            "project_code": project,
            "service": service,
            "person_id": person_id,
            "category_id": cat_id,
            "building_id": zone_id,
            "level_id": level_id,
            "area": area,
            "estimated_hours": round(row_hours, 2),
            "task_name_id": task_id,
            "task_sub_task_id": sub_id,
            "model_id": model_id,
            "drawing_name": dname,
            "drawing_number": dnum,
            "scale": scale,
            "paper_size": paper_size,
            "notes": _one(form, "notes"),
            "entered_by": entered_by,
            "created_at": dt.datetime.now().replace(microsecond=0).isoformat(sep=" "),
            "session_start": session_start,
        }

    con = connect(db_path)
    cur = con.cursor()
    model_names = {int(r["id"]): r["name"] for r in con.execute("SELECT id, name FROM models")}
    sub_names = {int(r["id"]): r["name"] for r in con.execute("SELECT id, name FROM task_sub_tasks")}
    name_args = (model_names, sub_names)
    if edit_id:
        cols = row_cols(
            dates[0], people[0], zone_id, level_id, model_id, work, hours,
            *name_args)
        keys = [k for k in cols if k not in ("entered_by", "created_at")]
        cur.execute(
            f"UPDATE forecast_assignments SET {','.join(k + '=?' for k in keys)} WHERE id=?",
            [cols[k] for k in keys] + [int(edit_id)],
        )
        con.commit()
        con.close()
        return 1, True

    combos = list(product(dates, people))
    if len(combos) > 2000:
        con.close()
        raise ValueError(
            f"That would create {len(combos)} rows ({len(people)} people × {len(dates)} days). "
            "Tick fewer days, or use Create loading plan for floor-by-floor loading "
            f"({len(people)} people × days on each floor).")
    keys = None
    n = 0
    assigned = 0.0
    last = len(combos) - 1
    each = round(hours / len(combos), 2) if combos else 0
    for i, (work_date, person_id) in enumerate(combos):
        row_hours = round(hours - assigned, 2) if i == last else each
        assigned += row_hours
        cols = row_cols(work_date, person_id, zone_id, level_id, model_id, work,
                        row_hours, *name_args)
        if keys is None:
            keys = list(cols.keys())
        cur.execute(
            f"INSERT INTO forecast_assignments({','.join(keys)}) VALUES ({','.join('?' * len(keys))})",
            [cols[k] for k in keys],
        )
        n += 1
    con.commit()
    con.close()
    return n, False


def delete_assignment(db_path, assignment_id):
    con = connect(db_path)
    con.execute("DELETE FROM forecast_assignments WHERE id=?", (int(assignment_id),))
    con.commit()
    con.close()


def working_dates(start, count, holidays=None):
    """count working days from start (inclusive). Weekly offs and public holidays skipped."""
    d = dt.date.fromisoformat(start)
    out = []
    guard = 0
    while len(out) < count and guard < 400:
        guard += 1
        if not is_leave_day(d, holidays):
            out.append(d.isoformat())
        d += dt.timedelta(days=1)
    return out


def next_open_dates(start, count, people, existing, holidays=None):
    """Next working dates from start that are not already fully booked for this team.

    Same-session saves skip days already used and continue. A new week's start
    date has no bookings, so it does not stick to last week's plan.
    """
    d = dt.date.fromisoformat(start)
    out = []
    guard = 0
    people = list(people)
    while len(out) < count and guard < 800:
        guard += 1
        iso = d.isoformat()
        busy = bool(people) and all((p, iso) in existing for p in people)
        if not is_leave_day(d, holidays) and not busy:
            out.append(iso)
        d += dt.timedelta(days=1)
    return out


def save_loading_plan(db_path, form, entered_by):
    """One floor at a time: every ticked person, then the next floor.

    Each floor has its own working-day count (3 was only an example).
    This is the management loading-sheet shape — one person-day at 8.25 hrs.
    """
    ensure_forecast_schema(db_path)
    project = _one(form, "project")
    if not project:
        raise ValueError("Choose a project.")
    start = _one(form, "plan_start") or _one(form, "start")
    if not start:
        raise ValueError("Choose a start date.")
    dt.date.fromisoformat(start)
    session_start = _one(form, "session_start") or _one(form, "start") or start
    try:
        dt.date.fromisoformat(session_start)
    except ValueError:
        session_start = start
    people = _ids(form, "person_ids")
    if not people:
        raise ValueError("Tick the people on this loading plan.")
    days_by_floor = {}
    for key, values in form.items():
        if not str(key).startswith("floor_days_"):
            continue
        try:
            lid = int(str(key)[11:])
        except ValueError:
            continue
        raw = (values[0] if values else "").strip()
        if not raw:
            continue
        try:
            n = int(raw)
        except ValueError:
            continue
        if n > 0:
            days_by_floor[lid] = max(1, min(30, n))
    order = []
    for bit in (_one(form, "floor_order") or "").split(","):
        bit = bit.strip()
        if bit.isdigit():
            order.append(int(bit))
    floors = [(lid, days_by_floor[lid]) for lid in order if lid in days_by_floor]
    if not floors:
        floors = list(days_by_floor.items())
    if not floors:
        raise ValueError("Set working days on at least one floor. Leave 0 to skip a floor.")
    zone_id = _int(form, "building_id")
    replace = (_one(form, "replace_session") or "").strip() in ("1", "yes", "on")
    append = (_one(form, "append_same_day") or "").strip() in ("1", "yes", "on")
    con = connect(db_path)
    allowed = _active_team_ids_db(con, project)
    if allowed is not None:
        people = [p for p in people if p in allowed]
        if not people:
            con.close()
            raise ValueError(
                "Those people are not on this project's team. "
                "Add them in Manage Lists → Project team.")
    prow = con.execute("SELECT no_zone FROM projects WHERE code=?", (project,)).fetchone()
    if prow and not prow["no_zone"] and not zone_id:
        all_z = con.execute(
            "SELECT id FROM buildings WHERE project_code=? AND is_all=1",
            (project,)).fetchone()
        zone_id = all_z["id"] if all_z else None
        if not zone_id:
            con.close()
            raise ValueError("Pick a zone (or add All zones in Manage Lists).")
    work_days = sum(n for _lid, n in floors)
    expected = len(people) * work_days
    if expected > 8000:
        one_each = len(people) * len(floors)
        con.close()
        raise ValueError(
            f"That would create {expected} rows ({len(people)} people × {work_days} floor-days). "
            f"A loading plan is people × days on each floor only — not models × sub-tasks × zones. "
            f"{len(people)} people × {len(floors)} floors × 1 day = {one_each} rows. "
            "Set days only on the floors you need (leave the rest 0).")
    cat_ids = _ids(form, "category_ids")
    if not cat_ids:
        one = _int(form, "category_id")
        cat_ids = [one] if one else []
    cat_id = cat_ids[0] if cat_ids else None
    model_id = _int(form, "model_id")
    task_id = None
    drawing_name = drawing_number = None
    if model_id:
        r = con.execute("SELECT name FROM models WHERE id=?", (model_id,)).fetchone()
        drawing_name = r["name"] if r else None
    if cat_ids:
        marks = ",".join("?" * len(cat_ids))
        names = [r["name"] for r in con.execute(
            f"SELECT name FROM categories WHERE id IN ({marks}) ORDER BY name",
            cat_ids,
        )]
        drawing_number = ", ".join(n for n in names if n) or None
    holidays = load_public_holiday_days(con=con)
    if replace and session_start:
        con.execute(
            "DELETE FROM forecast_assignments WHERE project_code=? AND session_start=?",
            (project, session_start),
        )
    booked = {
        (r["person_id"], r["work_date"])
        for r in con.execute(
            "SELECT person_id, work_date FROM forecast_assignments "
            "WHERE project_code=? AND work_date>=?",
            (project, start),
        )
    }
    booked_floor = {
        (r["person_id"], r["work_date"], r["level_id"])
        for r in con.execute(
            "SELECT person_id, work_date, level_id FROM forecast_assignments "
            "WHERE project_code=? AND work_date>=?",
            (project, start),
        )
    }
    now = dt.datetime.now().replace(microsecond=0).isoformat(sep=" ")
    n = 0
    skipped = 0
    cur = con.cursor()
    cursor = start
    all_dates = []
    for level_id, days_each in floors:
        if append:
            chunk = working_dates(cursor, days_each, holidays)
        else:
            chunk = next_open_dates(cursor, days_each, people, booked, holidays)
        if len(chunk) < days_each:
            con.close()
            raise ValueError("Could not fit that many working days from the start date.")
        all_dates.extend(chunk)
        for work_date in chunk:
            for person_id in people:
                if append:
                    if (person_id, work_date, level_id) in booked_floor:
                        skipped += 1
                        continue
                elif (person_id, work_date) in booked:
                    skipped += 1
                    continue
                cur.execute(
                    """INSERT INTO forecast_assignments(
                        work_date, project_code, service, person_id, category_id,
                        building_id, level_id, area, estimated_hours, task_name_id,
                        task_sub_task_id, model_id, drawing_name, drawing_number,
                        scale, paper_size, notes, entered_by, created_at, session_start)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (work_date, project, "Modelling Works", person_id, cat_id,
                     zone_id, level_id, None, HOURS_PER_DAY, task_id,
                     None, model_id, drawing_name, drawing_number,
                     None, None, "Loading plan — floor rotation", entered_by, now,
                     session_start),
                )
                booked.add((person_id, work_date))
                booked_floor.add((person_id, work_date, level_id))
                n += 1
        cursor = (dt.date.fromisoformat(chunk[-1]) + dt.timedelta(days=1)).isoformat()
    con.commit()
    con.close()
    return {
        "n": n,
        "skipped": skipped,
        "start": all_dates[0] if all_dates else start,
        "end": all_dates[-1] if all_dates else start,
        "days": len(all_dates) or 1,
        "floors": len(floors),
        "people": len(people),
        "work_days": work_days,
        "session_start": session_start,
        "replaced": replace,
        "appended": append,
    }


def _loading_plan_html(ref, project, start, hours_per_day, session_start=None,
                      floor_days=None, fill_days="", checked_people=None,
                      checked_cats=None, model_id=None, zone_id=None, replace=False,
                      append=False, week_qs=""):
    team_ids, any_team = current_team_ids(ref, project)
    if any_team:
        people = [p for p in (ref.get("people") or []) if str(p["id"]) in team_ids]
    else:
        people = list(ref.get("people") or [])
    if checked_people:
        have = {int(p["id"]) for p in people}
        extra = [p for p in (ref.get("people") or [])
                 if int(p["id"]) in {int(x) for x in checked_people} and int(p["id"]) not in have]
        people = people + extra
    levels = [l for l in (ref.get("levels") or []) if l.get("project_code") == project]
    zones = [b for b in (ref.get("buildings") or []) if b.get("project_code") == project]
    no_zone = False
    for p in ref.get("projects") or []:
        if p["code"] == project:
            no_zone = bool(p.get("no_zone"))
            break
    session_start = session_start or start
    want_people = None if checked_people is None else {int(x) for x in checked_people}
    want_cats = {int(x) for x in (checked_cats or [])}
    people_ticks = "".join(
        f'<label><input type="checkbox" name="person_ids" value="{p["id"]}"'
        f'{" checked" if (want_people is None or int(p["id"]) in want_people) else ""}>'
        f'{html.escape(p["name"])}'
        f'{(" (" + html.escape(p["emp_code"]) + ")") if p.get("emp_code") else ""}'
        f'</label>'
        for p in people
    ) or '<span class="none">No project team yet.</span>'
    models = [m for m in (ref.get("models") or []) if m.get("project_code") == project]
    model_opts = _opts(models, "id", "name",
                       selected=str(model_id or ""),
                       placeholder="model (Drawing Title)")
    cats = list(ref.get("categories") or [])
    cat_ticks = "".join(
        f'<label><input type="checkbox" name="category_ids" value="{c["id"]}"'
        f'{" checked" if int(c["id"]) in want_cats else ""}>'
        f'{html.escape(c["name"])}</label>'
        for c in cats
    ) or '<span class="none">No categories yet.</span>'
    floor_days = floor_days or {}
    floor_order = ",".join(str(l["id"]) for l in levels)
    floor_rows = "".join(
        '<div class="frow">'
        f'<span class="nm">{html.escape(l["label"])}</span>'
        f'<input type="number" name="floor_days_{l["id"]}" class="floor_days" min="0" max="30" '
        f'value="{html.escape(str(floor_days[l["id"]])) if l["id"] in floor_days else ""}" '
        f'placeholder="0" title="Working days on {html.escape(l["label"], quote=True)}">'
        '</div>'
        for l in levels
    ) or '<span class="none">No floors on this project.</span>'
    zone_field = ""
    if not no_zone:
        opts = []
        for z in zones:
            if zone_id:
                sel = " selected" if str(z["id"]) == str(zone_id) else ""
            else:
                sel = " selected" if z.get("is_all") else ""
            opts.append(f'<option value="{z["id"]}"{sel}>{html.escape(z["name"])}</option>')
        zone_field = (
            f'<div><label>Zone</label><select name="building_id">'
            f'{"".join(opts) or "<option value=\"\">(none)</option>"}</select></div>'
        )
    session_start = session_start or start
    if replace:
        heading = "Edit this plan"
        save_label = "Save changes"
        confirm_verb = "Replace this plan with"
        hint = (
            "These are the values saved on this plan. Change anything that is wrong and "
            "<b>Save changes</b> to replace this week’s rows. To keep these rows and load "
            "more people or floors on the <b>same date</b>, use "
            "<b>Add another plan this day</b>."
        )
    elif append:
        heading = "Add another plan this day"
        save_label = "Add to this day"
        confirm_verb = "Add"
        hint = (
            "This <b>adds</b> more people, floors or categories on this date. "
            "Rows already saved stay. Same person on a new floor the same day is listed "
            "together on the management Excel."
        )
    else:
        heading = "Add floors"
        save_label = "Add to this plan"
        confirm_verb = "Add"
        hint = (
            "One row = <b>one person × one floor × one day</b>. "
            "22 people × 6 floors × 1 day = <b>132 rows</b> — not models × sub-tasks × zones. "
            "Type days only on the floors you are doing (leave the rest 0). "
            "<b>Fill</b> writes the same number on every floor, so clear unused floors back to 0. "
            "Next floors = add again. Same week = one Excel. Next week = new Excel."
        )
    extra_hidden = ""
    if replace:
        extra_hidden = '<input type="hidden" name="replace_session" value="1">'
    elif append:
        extra_hidden = '<input type="hidden" name="append_same_day" value="1">'
    extra_actions = ""
    if week_qs:
        if replace:
            extra_actions = (
                f'<a class="btn green" href="/forecast?{html.escape(week_qs)}&amp;add=1">'
                f'Add another plan this day</a>'
            )
        elif append:
            extra_actions = (
                f'<a class="btn grey" href="/forecast?{html.escape(week_qs)}">'
                f'Back to edit this plan</a>'
            )
    team_open = " open" if (replace or append) else ""
    team_summary = (
        f'Team ({len(people)}) — use All / None, then tick who should be on this plan'
    )
    return f"""
  <form class="card" method="POST" action="/forecast/loading-plan" id="loading_plan">
    <input type="hidden" name="project" value="{html.escape(project)}">
    <input type="hidden" name="session_start" value="{html.escape(session_start)}">
    <input type="hidden" name="floor_order" value="{html.escape(floor_order)}">
    {extra_hidden}
    <div class="actions" style="margin:0 0 12px">
      <h2 style="margin:0;flex:1">{heading}</h2>
      {extra_actions}
    </div>
    <div class="steps">
      <span><b>1</b> floors</span>
      <span><b>2</b> add to this week</span>
      <span><b>3</b> more floors if needed</span>
      <span><b>4</b> export</span>
    </div>
    <p class="hint" style="margin-top:0">{hint}
       Each day is <b>{hours_per_day:g} hrs</b>.</p>
    <div class="grid">
      <div><label>Start date</label>
           <input type="date" name="plan_start" value="{html.escape(start)}" required></div>
      {zone_field}
      <div><label>Same days on every floor</label>
           <div style="display:flex;gap:8px;align-items:end">
             <input type="number" id="fill_floor_days" name="fill_floor_days" min="0" max="30"
                    value="{html.escape(fill_days)}" placeholder="days">
             <button type="button" id="fill_floors" class="btn grey">Fill</button>
           </div></div>
    </div>
    <div style="margin-top:14px">
      <label>Category
        <span class="bulk"><button type="button" data-tick-all="load_cats">All</button>
        <button type="button" data-tick-none="load_cats">None</button></span></label>
      <div class="ticks" id="load_cats">{cat_ticks}</div>
      <p class="hint" style="margin:6px 0 0">Ticked categories go to <b>Dwg_number</b> on the
         management sheet. Sub-tasks are not reported there.</p>
    </div>
    <div style="margin-top:14px">
      <label>Model</label>
      <select name="model_id">{model_opts}</select>
      <p class="hint" style="margin:6px 0 0">Goes to <b>Drawing Title</b> on the management sheet.</p>
    </div>
    <div style="margin-top:14px"><label>Days on each floor</label>
        <div class="floordays" id="load_floors">{floor_rows}</div></div>
    <details class="teamfold"{team_open}>
      <summary>{html.escape(team_summary)}</summary>
      <div style="margin-top:8px">
        <span class="bulk"><button type="button" data-tick-all="load_people">All</button>
        <button type="button" data-tick-none="load_people">None</button></span>
      </div>
      <div class="people ticks" id="load_people" style="margin-top:8px">{people_ticks}</div>
    </details>
    <div class="countbar" id="load_count">Set days on a floor to see the plan.</div>
    <div class="actions">
      <button type="submit">{save_label}</button>
      {extra_actions}
    </div>
  </form>
  <script>
  (function(){{
    var form=document.getElementById("loading_plan");
    if(!form) return;
    var bar=document.getElementById("load_count");
    var hrs={hours_per_day};
    function nPeople(){{
      var b=document.getElementById("load_people"); if(!b) return 0;
      return b.querySelectorAll("input:checked").length;
    }}
    function floorDays(){{
      return [].slice.call(form.querySelectorAll(".floor_days")).map(function(el){{
        var n=parseInt(el.value||"0",10); return isNaN(n)||n<0?0:n;
      }});
    }}
    function preview(){{
      if(!bar) return;
      var p=nPeople(), days=floorDays(), used=days.filter(function(n){{return n>0;}});
      var f=used.length;
      var d=days.reduce(function(a,b){{return a+b;}},0);
      if(!p || !d){{
        bar.innerHTML = "Type days on the floors you are loading now. Leave the rest 0.";
        return;
      }}
      var same = used.every(function(n){{ return n===used[0]; }});
      var how = same
        ? (p+" people × "+f+" floor"+(f===1?"":"s")+" × "+used[0]+" day"+(used[0]===1?"":"s")+" each")
        : (p+" people × "+d+" floor-days");
      bar.innerHTML = how+" = <b>"+(p*d)+"</b> person-days. Add, then export.";
    }}
    form.addEventListener("change", preview);
    form.addEventListener("input", preview);
    preview();
    var fillBtn=document.getElementById("fill_floors");
    if(fillBtn) fillBtn.addEventListener("click", function(){{
      var n=(document.getElementById("fill_floor_days")||{{}}).value;
      [].forEach.call(form.querySelectorAll(".floor_days"), function(el){{ el.value=n; }});
      preview();
    }});
    form.addEventListener("submit", function(ev){{
      var p=nPeople(), days=floorDays(), d=days.reduce(function(a,b){{return a+b;}},0);
      var f=days.filter(function(n){{return n>0;}}).length;
      if(!p||!d){{ ev.preventDefault(); alert("Set working days on at least one floor."); return; }}
      if(!confirm("{confirm_verb} "+(p*d)+" rows ("+p+" people × "+d+" floor-days on "+f+" floor(s))?"))
        ev.preventDefault();
    }});
  }})();
  </script>
"""


def _export_rows(db_path, project, start, days, service, person_id=None):
    end = window_end(start, days)
    con = connect(db_path)
    sql = """
        SELECT f.service, f.work_date, c.name AS category, b.name AS zone, l.label AS level,
               f.area, f.estimated_hours, pe.emp_code, ts.name AS sub_task, tn.name AS task,
               m.name AS model, f.drawing_name, f.drawing_number, f.scale, f.paper_size
        FROM forecast_assignments f
        JOIN people pe ON pe.id = f.person_id
        LEFT JOIN categories c ON c.id = f.category_id
        LEFT JOIN buildings b ON b.id = f.building_id
        LEFT JOIN levels l ON l.id = f.level_id
        LEFT JOIN task_names tn ON tn.id = f.task_name_id
        LEFT JOIN task_sub_tasks ts ON ts.id = f.task_sub_task_id
        LEFT JOIN models m ON m.id = f.model_id
        WHERE f.project_code = ? AND f.service = ?
          AND """ + session_scope_sql("f")
    params = [project, service, start, start, end]
    if person_id:
        sql += " AND f.person_id=?"
        params.append(person_id)
    sql += " ORDER BY f.work_date, pe.emp_code, f.id"
    rows = con.execute(sql, params).fetchall()
    no_zone = con.execute(
        "SELECT no_zone FROM projects WHERE code=?", (project,)
    ).fetchone()
    con.close()
    hide_zone = bool(no_zone and no_zone[0])
    return [dict(r) for r in rows], hide_zone


def _export_compare(db_path, project, start, days, person_id=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    cmp = compare_forecast_actual(db_path, project, start, days, person_id)
    fname = f"forecast_vs_actual_{project}_{start}_{window_end(start, days)}.xlsx"
    wb = Workbook()
    thin = Border(
        left=Side(style="thin", color="BBBBBB"),
        right=Side(style="thin", color="BBBBBB"),
        top=Side(style="thin", color="BBBBBB"),
        bottom=Side(style="thin", color="BBBBBB"),
    )
    fill = PatternFill("solid", fgColor=INK)
    hdr_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=16, color=INK)
    sub_font = Font(size=10, color="555555")

    def write_sheet(ws, title, headers, rows, widths):
        ncol = len(headers)
        ws["A1"] = title
        ws["A1"].font = title_font
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        ws["A2"] = f"project {project}  ·  {start} to {window_end(start, days)}"
        ws["A2"].font = sub_font
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
        for i, h in enumerate(headers, 1):
            c = ws.cell(3, i, h)
            c.font = hdr_font
            c.fill = fill
            c.alignment = Alignment(vertical="center")
            c.border = thin
        r = 4
        for row in rows:
            for i, v in enumerate(row, 1):
                c = ws.cell(r, i, v)
                c.border = thin
            r += 1
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(ncol)}{max(3, r - 1)}"

    ws = wb.active
    ws.title = "By person"
    person_rows = [
        [r["person"], r["emp_code"], r["forecast"], r["actual"], r["diff"]]
        for r in cmp["people"]
    ]
    write_sheet(ws, "FORECAST VS ACTUAL — HOURS",
                ["Person", "Emp ID", "Forecast hrs", "Actual hrs", "Actual − Forecast"],
                person_rows, [28, 12, 14, 14, 18])

    ws2 = wb.create_sheet("Planned vs logged")
    item_rows = [
        [r["work_date"], r["person"], r.get("category") or "", r.get("task") or "",
         r.get("sub_task") or "", r["forecast"], "Yes" if r.get("logged") else "No"]
        for r in cmp["items"]
    ]
    write_sheet(ws2, "FORECAST VS ACTUAL — WORK ITEMS",
                ["Date", "Person", "Category", "Task", "Sub Task", "Forecast hrs", "Logged"],
                item_rows, [12, 28, 16, 22, 32, 14, 10])

    buf = io.BytesIO()
    wb.save(buf)
    return fname, buf.getvalue()


def _export_progress(db_path, project):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    prog = project_progress(db_path, project)
    fname = f"project_progress_{project}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Progress"
    thin = Border(
        left=Side(style="thin", color="BBBBBB"),
        right=Side(style="thin", color="BBBBBB"),
        top=Side(style="thin", color="BBBBBB"),
        bottom=Side(style="thin", color="BBBBBB"),
    )
    fill = PatternFill("solid", fgColor=INK)
    hdr_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=16, color=INK)
    total = prog["total"] or 0
    ws["A1"] = "PROJECT PROGRESS — HOURS TO COMPLETE"
    ws["A1"].font = title_font
    ws.merge_cells("A1:F1")
    ws["A2"] = (
        f"Project {project}  ·  Completed {_pct(prog['completed'], total):g}%  ·  "
        f"Ongoing {_pct(prog['ongoing'], total):g}%  ·  "
        f"Not started {_pct(prog['not_started'], total):g}%"
    )
    headers = ["Category", "Sub Task", "Zone", "Level", "Hours to complete", "Latest status", "Logged hrs"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(4, i, h)
        c.font = hdr_font
        c.fill = fill
        c.border = thin
    r = 5
    for row in prog["rows"]:
        vals = [row["category"], row["sub_task"], row.get("zone") or "", row.get("level") or "",
                row["hours"], row["status"], row["logged"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.border = thin
        r += 1
    for i, w in enumerate([16, 36, 14, 12, 18, 16, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return fname, buf.getvalue()


LOADING_HEADERS = [
    "Service", "Floor", "Dwg_number", "Drawing Title", "project_scope", "stage",
    "Paper Size", "Scale", "Engest Hrs", "Drawing Status", "Due date", "Priority", "Eng_ID",
]


def _export_loading(db_path, project, start, days, person_id=None):
    """Management loading sheet. One row per person-day (8.25 hrs). No sub-tasks."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    end = window_end(start, days)
    con = connect(db_path)
    prj = con.execute(
        "SELECT code, name, internal_code, hours_per_day FROM projects WHERE code=?", (project,)
    ).fetchone()
    day_hours = float(prj["hours_per_day"]) if prj and prj["hours_per_day"] else HOURS_PER_DAY
    load_params = [project, start, start, end]
    if person_id:
        load_sql_tail = " AND f.person_id=?"
        load_params.append(person_id)
    else:
        load_sql_tail = ""
    rows = con.execute(
        """
        SELECT f.service, f.work_date, f.person_id,
               pe.emp_code,
               l.label AS level,
               c.name AS category,
               tn.name AS task,
               ts.name AS sub_task,
               m.name AS model,
               f.drawing_name, f.drawing_number,
               (SELECT mt.name FROM master_tasks mt
                  JOIN category_tasks ct ON ct.master_task_id = mt.id
                 WHERE mt.project_code = f.project_code AND ct.category_id = f.category_id
                 LIMIT 1) AS master_task
          FROM forecast_assignments f
          JOIN people pe ON pe.id = f.person_id
          LEFT JOIN categories c ON c.id = f.category_id
          LEFT JOIN levels l ON l.id = f.level_id
          LEFT JOIN task_names tn ON tn.id = f.task_name_id
          LEFT JOIN task_sub_tasks ts ON ts.id = f.task_sub_task_id
          LEFT JOIN models m ON m.id = f.model_id
         WHERE f.project_code = ? AND """ + session_scope_sql("f") + load_sql_tail + """
         ORDER BY f.work_date, pe.emp_code, f.id
        """,
        load_params,
    ).fetchall()
    status_sql = """
        SELECT ta.person_id, t.work_date, COALESCE(st.name, '') AS status
          FROM tasks t
          JOIN task_assignees ta ON ta.task_id = t.id
          LEFT JOIN statuses st ON st.id = t.status_id
         WHERE t.project_code = ? AND t.work_date >= ? AND t.work_date <= ?
        """
    status_params = [project, start, end]
    if person_id:
        status_sql += " AND ta.person_id=?"
        status_params.append(person_id)
    status_sql += " ORDER BY t.work_date, t.id"
    status_rows = con.execute(status_sql, status_params).fetchall()
    con.close()
    latest_status = {}
    for r in status_rows:
        latest_status[(r["person_id"], r["work_date"])] = r["status"]

    days_out = []
    index = {}
    for row in rows:
        key = (row["person_id"], row["work_date"])
        if key not in index:
            emp = row["emp_code"] or ""
            try:
                emp_out = int(emp) if str(emp).isdigit() else emp
            except (TypeError, ValueError):
                emp_out = emp
            rec = {
                "services": [], "levels": [], "dwgs": [], "titles": [], "scopes": [],
                "status": latest_status.get(key, "") or "",
                "work_date": row["work_date"] or "",
                "emp": emp_out,
            }
            index[key] = rec
            days_out.append(rec)
        rec = index[key]
        stored_dwg = (row["drawing_number"] or "").strip()
        sub = (row["sub_task"] or "").strip()
        # Dwg_number = selected category. Ignore leftover sub-task names.
        if stored_dwg and stored_dwg != sub:
            dwg = stored_dwg
        else:
            dwg = (row["category"] or "").strip()
        title = (row["model"] or row["drawing_name"] or "").strip()
        for lst, val in (
            (rec["services"], row["service"]),
            (rec["levels"], row["level"]),
            (rec["dwgs"], dwg),
            (rec["titles"], title),
            (rec["scopes"], row["master_task"] or row["category"]),
        ):
            s = (val or "").strip()
            if s and s not in lst:
                lst.append(s)

    internal = (prj["internal_code"] if prj and prj["internal_code"] else project) if prj else project
    fname = f"loading_task_{internal}_{project}_{start}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    thin = Border(
        left=Side(style="thin", color="BBBBBB"),
        right=Side(style="thin", color="BBBBBB"),
        top=Side(style="thin", color="BBBBBB"),
        bottom=Side(style="thin", color="BBBBBB"),
    )
    fill = PatternFill("solid", fgColor=INK)
    hdr_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14, color=INK)

    ws["A1"] = f"project {internal}"
    ws["A1"].font = title_font
    for i, h in enumerate(LOADING_HEADERS, 1):
        c = ws.cell(2, i, h)
        c.font = hdr_font
        c.fill = fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = thin
    r = 3
    for row in days_out:
        vals = [
            ", ".join(row["services"]),
            ", ".join(row["levels"]),
            ", ".join(row["dwgs"]),
            ", ".join(row["titles"]),
            ", ".join(row["scopes"]),
            "LOD 500",
            "A0",
            "N.A",
            day_hours,
            row["status"],
            row["work_date"],
            "",
            row["emp"],
        ]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.border = thin
            if i == 9:
                c.number_format = "0.00"
        r += 1
    widths = [16, 12, 28, 28, 28, 12, 12, 10, 12, 16, 12, 10, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:M{max(2, r - 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    return fname, buf.getvalue()


def export_report(db_path, project, start, days, kind, person_id=None):
    """Return (filename, xlsx_bytes) for modelling, sheet, compare, progress, or loading."""
    ensure_forecast_schema(db_path)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if kind == "compare":
        return _export_compare(db_path, project, start, days, person_id)
    if kind == "progress":
        return _export_progress(db_path, project)
    if kind == "loading":
        start, days, _end = plan_window_or_latest(db_path, project, start, days)
        return _export_loading(db_path, project, start, days, person_id)

    if kind == "sheet":
        service = "Sheet Work"
        sheet_name = "Sheet Work"
        base_headers = ["Service", "Work Date", "Category"]
        extra = ["Drawing Name", "Drawing Number", "Scale", "Paper Size"]
        fname = f"forecast_sheet_{project}_{start}_{window_end(start, days)}.xlsx"
    else:
        service = "Modelling Works"
        sheet_name = "Modelling Works"
        base_headers = ["Service", "Work Date", "Category"]
        extra = ["Model Name"]
        fname = f"forecast_modelling_{project}_{start}_{window_end(start, days)}.xlsx"

    rows, hide_zone = _export_rows(db_path, project, start, days, service, person_id)
    headers = list(base_headers)
    if not hide_zone:
        headers.append("Zone")
    headers += ["Level", "Area", "Estimated Hours", "Emp ID", "Sub Task", "Task"] + extra

    def cells(r):
        out = [r["service"], r["work_date"], r["category"] or ""]
        if not hide_zone:
            out.append(r["zone"] or "")
        out += [
            r["level"] or "",
            r["area"] if r["area"] is not None else "",
            r["estimated_hours"],
            r["emp_code"] or "",
            r["sub_task"] or "",
            r["task"] or "",
        ]
        if kind == "sheet":
            out += [
                r["drawing_name"] or "",
                r["drawing_number"] or "",
                r["scale"] or "",
                r["paper_size"] or "",
            ]
        else:
            out.append(r["model"] or "")
        return out

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    thin = Border(
        left=Side(style="thin", color="BBBBBB"),
        right=Side(style="thin", color="BBBBBB"),
        top=Side(style="thin", color="BBBBBB"),
        bottom=Side(style="thin", color="BBBBBB"),
    )
    fill = PatternFill("solid", fgColor=INK)
    hdr_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=16, color=INK)
    sub_font = Font(size=10, color="555555")
    ncol = len(headers)

    ws["A1"] = f"project {project}"
    ws["A1"].font = title_font
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws["A2"] = f"{service}  ·  {start} to {window_end(start, days)}"
    ws["A2"].font = sub_font
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)

    for i, h in enumerate(headers, 1):
        c = ws.cell(3, i, h)
        c.font = hdr_font
        c.fill = fill
        c.border = thin
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for r_i, row in enumerate(rows, 4):
        for c_i, v in enumerate(cells(row), 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = thin
            cell.alignment = Alignment(vertical="top")
            if headers[c_i - 1] in ("Area", "Estimated Hours") and v != "":
                cell.number_format = "0.00"

    widths = {
        "Service": 18, "Work Date": 12, "Category": 14, "Zone": 14, "Level": 12,
        "Area": 10, "Estimated Hours": 16, "Emp ID": 12, "Sub Task": 28, "Task": 22,
        "Model Name": 36, "Drawing Name": 32, "Drawing Number": 16, "Scale": 10,
        "Paper Size": 12,
    }
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 16)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(ncol)}{max(3, 3 + len(rows))}"

    buf = io.BytesIO()
    wb.save(buf)
    return fname, buf.getvalue()


def _opts(items, value_key, label_key, selected="", placeholder="choose"):
    s = f'<option value="">{html.escape("— " + placeholder + " —")}</option>'
    for it in items:
        val = str(it[value_key])
        lab = it[label_key]
        sel = " selected" if selected != "" and str(selected) == val else ""
        s += f'<option value="{html.escape(val)}"{sel}>{html.escape(str(lab))}</option>'
    return s


def _load_grid_html(ref, project, dates, rows, hours_per_day, can_write, person_id=None):
    team_ids, any_team = current_team_ids(ref, project)
    people = list(ref.get("people") or [])
    if person_id:
        people = [p for p in people if str(p["id"]) == str(person_id)]
        if any_team:
            people = [p for p in people if str(p["id"]) in team_ids]
    elif any_team:
        people = [p for p in people if str(p["id"]) in team_ids]
    if not people:
        return '<p class="sub">No people on this project team yet — add them in Manage Lists.</p>'

    hours = {}
    for r in rows:
        key = (str(r["person_id"]), r["work_date"])
        hours[key] = hours.get(key, 0.0) + float(r["estimated_hours"] or 0)

    head = "<th>Person</th><th>Emp ID</th>" + "".join(
        f'<th class="num">{html.escape(d[5:])}</th>' for d in dates
    ) + '<th class="num">Total</th>'
    body = []
    for p in people:
        pid = str(p["id"])
        if any_team and pid not in team_ids:
            continue
        cells = []
        total = 0.0
        for d in dates:
            h = hours.get((pid, d), 0.0)
            total += h
            cls = "empty" if h == 0 else ("over" if h > hours_per_day + 0.001 else "okh")
            label = f"{h:g}" if h else "—"
            if can_write:
                inner = (f'<button type="button" class="cellbtn" data-date="{html.escape(d)}" '
                         f'data-person="{html.escape(pid)}">{label}</button>')
            else:
                inner = label
            cells.append(f'<td class="num {cls}">{inner}</td>')
        emp = html.escape(p.get("emp_code") or "")
        tcls = "over" if total > hours_per_day * len(dates) + 0.001 else "num"
        body.append(
            "<tr>"
            f'<td>{html.escape(p["name"])}</td><td>{emp}</td>'
            + "".join(cells)
            + f'<td class="num {tcls}"><b>{total:g}</b></td></tr>'
        )
    cap = f"{hours_per_day:g} h/day"
    return (
        f'<p class="hint">Hours assigned vs project day length ({html.escape(cap)}). '
        f'Red = over the day. Click a cell to assign that person on that date.</p>'
        f'<div class="scroll"><table><thead><tr>{head}</tr></thead>'
        f'<tbody>{"".join(body) or "<tr><td colspan=99 class=hint>No team members.</td></tr>"}'
        f'</tbody></table></div>'
    )


def _assignment_table(rows, show_zone, can_write, qs):
    zone_h = "<th>Zone</th>" if show_zone else ""
    extra_h = "<th>Model / Drawing</th>"
    act_h = "<th class='act'>Edit</th><th class='act'>Delete</th>" if can_write else ""
    head = (f"{act_h}<th>Date</th><th>Service</th><th>Person</th><th>Emp ID</th><th>Category</th>"
            f"{zone_h}<th>Level</th><th>Task</th><th>Sub Task</th><th class='num'>Hours</th>"
            f"{extra_h}")
    body = []
    for r in rows:
        extra = r.get("model") or ""
        if r["service"] == "Sheet Work":
            bits = [r.get("drawing_name") or "", r.get("drawing_number") or "",
                    r.get("scale") or "", r.get("paper_size") or ""]
            extra = " · ".join(b for b in bits if b)
        zone_c = f'<td>{html.escape(r.get("zone") or "")}</td>' if show_zone else ""
        actions = ""
        if can_write:
            actions = (
                f'<td class="act"><a class="btn sm" href="/forecast?{qs}&edit={r["id"]}">Edit</a></td>'
                f'<td class="act"><form method="POST" action="/forecast/delete" style="margin:0" '
                f"onsubmit=\"return confirm('Delete this assignment?')\">"
                f'<input type="hidden" name="id" value="{r["id"]}">'
                f'<input type="hidden" name="project" value="{html.escape(r.get("project_code") or "")}">'
                f'<input type="hidden" name="qs" value="{html.escape(qs)}">'
                f'<button type="submit" class="danger sm">Delete</button></form></td>'
            )
        body.append(
            "<tr>"
            f"{actions}"
            f'<td>{html.escape(str(r["work_date"]))}</td>'
            f'<td>{html.escape(r["service"])}</td>'
            f'<td>{html.escape(r.get("person") or "")}</td>'
            f'<td>{html.escape(r.get("emp_code") or "")}</td>'
            f'<td>{html.escape(r.get("category") or "")}</td>'
            f"{zone_c}"
            f'<td>{html.escape(r.get("level") or "")}</td>'
            f'<td>{html.escape(r.get("task") or "")}</td>'
            f'<td>{html.escape(r.get("sub_task") or "")}</td>'
            f'<td class="num">{r["estimated_hours"]:g}</td>'
            f'<td>{html.escape(extra)}</td></tr>'
        )
    if not body:
        span = 11 + (1 if show_zone else 0) + (2 if can_write else 0)
        body.append(
            f'<tr><td colspan="{span}" class="hint">No assignments in this date range. '
            'Open a saved period above, then Export loading sheet.</td></tr>')
    return f'<div class="scroll"><table><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'


def _tick_block(title, box_id, filter_ph=""):
    filt = (f'<input class="filter" data-filter="{box_id}" placeholder="{html.escape(filter_ph)}">'
            if filter_ph else "")
    return (
        f'<div class="tickcol" id="{box_id}_wrap">'
        f'<label>{html.escape(title)}'
        f'<span class="bulk"><button type="button" data-tick-all="{box_id}">All</button>'
        f'<button type="button" data-tick-none="{box_id}">None</button></span></label>'
        f'{filt}<div class="ticks" id="{box_id}"></div></div>'
    )


def _h(n):
    return f"{n:g}" if abs(n) >= 0.005 else "0"


def _var_cls(diff):
    if diff > 0.05:
        return "okh"
    if diff < -0.05:
        return "short"
    return "num"


def _periods_html(periods, project, start, days, can_write=False, add_href=""):
    if not periods:
        return ""
    chips = ['<span class="lab">Saved weeks</span>']
    for p in periods:
        qs = window_qs(project, p["start"], p["days"], end=p.get("end") or "")
        on = "on" if p["start"] == start and p["days"] == days else ""
        link = (
            f'<a class="{on}" href="/forecast?{qs}">'
            f'{html.escape(p["start"])} → {html.escape(p["end"])}'
            f' <span class="hint">({p["rows"]} rows)</span></a>'
        )
        rm = ""
        if can_write:
            sess = p.get("session_start") or ""
            qname = html.escape(f'{p["start"]} to {p["end"]}', quote=True)
            rm = (
                f'<form method="POST" action="/forecast/delete-plan" '
                f"onsubmit=\"return confirm('Remove the loading plan {qname}? "
                f"Daily logged tasks stay. This only deletes the forecast rows.')\">"
                f'<input type="hidden" name="project" value="{html.escape(project)}">'
                f'<input type="hidden" name="session_start" value="{html.escape(sess)}">'
                f'<input type="hidden" name="start" value="{html.escape(p["start"])}">'
                f'<input type="hidden" name="end" value="{html.escape(p["end"])}">'
                f'<input type="hidden" name="qs" value="{html.escape(qs)}">'
                f'<button type="submit" class="rm" title="Remove this plan">Remove</button></form>'
            )
        chips.append(f'<span class="planchip">{link}{rm}</span>')
    if add_href:
        chips.append(
            f'<a class="btn green" href="{html.escape(add_href)}" style="margin-left:6px">'
            f'Add another plan this day</a>'
        )
    return f'<div class="periods">{"".join(chips)}</div>'


def _compare_html(cmp, qs, table_class=""):
    people = cmp.get("people") or []
    items = cmp.get("items") or []
    tbl = f'table class="{html.escape(table_class)}"' if table_class else "table"
    if not people and not items:
        return ('<p class="hint">No forecast in this window yet — save assignments, '
                "then come back after the team logs daily tasks to compare.</p>")
    body = []
    tf = ta = 0.0
    for r in people:
        tf += r["forecast"]
        ta += r["actual"]
        body.append(
            "<tr>"
            f'<td>{html.escape(r["person"])}</td>'
            f'<td>{html.escape(r["emp_code"])}</td>'
            f'<td class="num">{_h(r["forecast"])}</td>'
            f'<td class="num">{_h(r["actual"])}</td>'
            f'<td class="num {_var_cls(r["diff"])}">{_h(r["diff"])}</td></tr>'
        )
    body.append(
        "<tr class='total'><td><b>Total</b></td><td></td>"
        f'<td class="num"><b>{_h(tf)}</b></td>'
        f'<td class="num"><b>{_h(ta)}</b></td>'
        f'<td class="num {_var_cls(ta - tf)}"><b>{_h(ta - tf)}</b></td></tr>'
    )
    person_tbl = (
        f'<div class="scroll"><{tbl}><thead><tr>'
        "<th>Person</th><th>Emp ID</th><th class='num'>Forecast hrs</th>"
        "<th class='num'>Actual hrs</th><th class='num'>Actual − Forecast</th>"
        f"</tr></thead><tbody>{''.join(body)}</tbody></table></div>"
    )
    ibody = []
    for r in items:
        logged = "Yes" if r.get("logged") else "No"
        lcls = "okh" if r.get("logged") else "short"
        extra = " <span class='hint'>(logged, not in forecast)</span>" if r.get("extra") else ""
        ibody.append(
            "<tr>"
            f'<td>{html.escape(str(r["work_date"]))}</td>'
            f'<td>{html.escape(r["person"])}</td>'
            f'<td>{html.escape(r.get("category") or "")}</td>'
            f'<td>{html.escape(r.get("task") or "")}</td>'
            f'<td>{html.escape(r.get("sub_task") or "")}{extra}</td>'
            f'<td class="num">{_h(r["forecast"])}</td>'
            f'<td class="{lcls}">{logged}</td></tr>'
        )
    item_tbl = ""
    if ibody:
        item_tbl = (
            '<h2>Planned work vs logged</h2>'
            '<p class="hint">Yes = that person logged the same sub-task on the daily form that day.</p>'
            f'<div class="scroll"><{tbl}><thead><tr>'
            "<th>Date</th><th>Person</th><th>Category</th><th>Task</th><th>Sub Task</th>"
            "<th class='num'>Forecast hrs</th><th>Logged</th>"
            f"</tr></thead><tbody>{''.join(ibody)}</tbody></table></div>"
        )
    return (
        f'<p class="hint">Forecast hours stay saved. Actual hours come from Daily Task Entry '
        f'({html.escape(cmp["start"])} to {html.escape(cmp["end"])}). '
        f'Amber = less done than planned. Green = more hours logged than planned.</p>'
        f"{person_tbl}{item_tbl}"
        f'<div class="exports" style="margin-top:10px">'
        f'<a class="btn green" href="/forecast/export?{qs}&kind=compare">Export Forecast vs Actual</a>'
        f"</div>"
    )


def render_forecast(db_path, q, banner, can_write, ref, flash=""):
    ensure_forecast_schema(db_path)
    project, start, days = window_from(q, ref, db_path)
    team_set = project_team_configured(ref, project) if project else False
    ref_page = without_former_members(ref)
    person_tok = ""
    person_id = None
    dates = date_list(start, days) if project else []
    end = window_end(start, days) if dates else ""
    qs = window_qs(project, start, days, end=end) if project else ""
    proj = project_row(ref, project)
    hours_per_day = float(proj["hours_per_day"]) if proj and proj.get("hours_per_day") else HOURS_PER_DAY
    default_total = hours_per_day * max(1, days)
    show_zone = not (proj and proj.get("no_zone"))
    rows = list_assignments(db_path, project, start, days, person_id) if project else []
    exp_start, exp_days, exp_end = (
        plan_window_or_latest(db_path, project, start, days) if project else (start, days, end)
    )
    load_qs = window_qs(project, exp_start, exp_days, end=exp_end) if project else ""

    edit = None
    edit_id = (q.get("edit", [""])[0] or "").strip()
    if edit_id and can_write:
        edit = get_assignment(db_path, edit_id)
        if edit and edit["project_code"] != project:
            project = edit["project_code"]
            team_set = project_team_configured(ref, project)
            proj = project_row(ref, project)
            show_zone = not (proj and proj.get("no_zone"))
            qs = window_qs(project, start, days, end=end)
            rows = list_assignments(db_path, project, start, days, person_id)

    ok = ""
    if flash:
        cls = "err" if flash.lower().startswith("error") or flash.lower().startswith("could") else "ok"
        ok = f'<div class="{cls}">{html.escape(flash)}</div>'
    elif "ok" in q:
        n = (q.get("n", [""])[0] or "").strip()
        if "upd" in q:
            ok = '<div class="ok">Updated assignment.</div>'
        elif n and n != "1":
            ok = f'<div class="ok">Saved {html.escape(n)} assignments.</div>'
        else:
            ok = '<div class="ok">Saved assignment.</div>'
    elif "del" in q:
        ok = '<div class="ok">Removed assignment.</div>'

    proj_opts = "".join(
        f'<option value="{html.escape(p["code"])}"{" selected" if p["code"]==project else ""}>'
        f'{html.escape(p["code"])} — {html.escape(p["name"])}</option>'
        for p in ref.get("projects") or []
    )
    day_dates = list(dates)
    if edit and edit["work_date"] not in day_dates:
        day_dates = [edit["work_date"]] + day_dates
    day_ticks = "".join(
        f'<label><input type="checkbox" name="work_dates" value="{html.escape(d)}">'
        f'{html.escape(d[5:])}</label>'
        for d in day_dates
    )
    svc_html = "".join(
        f'<label><input type="radio" name="service" value="{html.escape(s)}"'
        f'{" checked" if s == SERVICES[0] else ""}> {html.escape(s)}</label>'
        for s in SERVICES
    )
    scale_opts = '<option value="">— choose —</option>' + "".join(
        f'<option value="{html.escape(s)}">{html.escape(s)}</option>' for s in SCALES
    )
    paper_opts = '<option value="">— choose —</option>' + "".join(
        f'<option value="{html.escape(s)}">{html.escape(s)}</option>' for s in PAPER_SIZES
    )

    form_html = ""
    if can_write and project:
        eid = html.escape(str(edit["id"])) if edit else ""
        heading = f'Edit assignment #{eid}' if edit else "One person / one floor"
        save_label = "Update" if edit else "Save this row"
        day_word = "day" if days == 1 else "days"
        hrs_hint = ("— this row" if edit else
                    f"— {hours_per_day:g} hrs/day × days ticked, e.g. {default_total:g} for {days} {day_word}")
        del_btn = ""
        if edit:
            del_btn = (
                f'<form method="POST" action="/forecast/delete" style="display:inline" '
                f"onsubmit=\"return confirm('Delete this assignment?')\">"
                f'<input type="hidden" name="id" value="{eid}">'
                f'<input type="hidden" name="qs" value="{html.escape(qs)}">'
                f'<button type="submit" class="danger">Delete</button></form>'
            )
        form_html = f"""
  <form class="card" method="POST" action="/forecast/add" id="assign">
    <input type="hidden" name="project" id="as_project" value="{html.escape(project)}">
    <input type="hidden" name="session_start" value="{html.escape(start)}">
    <input type="hidden" name="edit_id" id="edit_id" value="{eid}">
    <input type="hidden" name="start" value="{html.escape(start)}">
    <input type="hidden" name="days" value="{days}">
    <input type="hidden" name="end" value="{html.escape(end)}">
    <input type="hidden" name="person" value="{html.escape(person_tok)}">
    <h2 style="margin-top:0">{html.escape(heading)}</h2>
    <p class="hint" style="margin-top:0">Use this only to fix one row. To load several floors,
       use <b>Add floors</b> above.</p>
    <div class="grid">
      <div class="full"><label>Service</label><div class="svc">{svc_html}</div></div>
      <div><label>{"Hours" if edit else "Total hours"} <span class="hint">{hrs_hint}</span></label>
           <input type="number" name="estimated_hours" id="estimated_hours" min="0.25" step="0.25"
                  value="{default_total:g}" required></div>
      <div><label>Area m²</label>
           <input type="number" name="area" id="area" min="0" step="0.01" placeholder="optional"></div>
      <div class="full"><label>Days
           <span class="bulk"><button type="button" data-tick-all="work_dates">All</button>
           <button type="button" data-tick-none="work_dates">None</button></span></label>
           <div class="days" id="work_dates">{day_ticks}</div></div>
    </div>
    <div class="grid3" style="margin-top:14px">
      {_tick_block("People", "people_ids", "Filter people")}
      {_tick_block("Category", "category_ids", "Filter category")}
      {_tick_block("Task", "task_name_ids")}
    </div>
    <div class="grid3" style="margin-top:14px">
      {_tick_block("Sub Task", "task_sub_task_ids", "Filter sub task")}
      <div id="zone_wrap">{_tick_block("Zone", "building_ids", "Filter zone")}</div>
      {_tick_block("Floor", "level_ids", "Filter floor")}
    </div>
    <div class="grid3" style="margin-top:14px">
      <div id="modelling_fields" class="full">{_tick_block("Model", "model_ids", "Filter model")}</div>
      <div id="sheet_fields_name"><label>Drawing name</label>
           <input name="drawing_name" id="drawing_name" placeholder="drawing title"></div>
      <div id="sheet_fields_num"><label>Drawing number</label>
           <input name="drawing_number" id="drawing_number"></div>
      <div id="sheet_fields_scale"><label>Scale</label><select name="scale" id="scale">{scale_opts}</select></div>
      <div id="sheet_fields_paper"><label>Paper size</label>
           <select name="paper_size" id="paper_size">{paper_opts}</select></div>
    </div>
    <div style="margin-top:14px"><label>Notes</label>
         <textarea name="notes" id="notes" rows="2"></textarea></div>
    <div class="countbar" id="combo_count">Tick people, days and floors to see how many rows will be created.</div>
    <div class="actions">
      <button type="submit">{save_label}</button>
      {del_btn}
      <a class="btn grey" href="/forecast?{html.escape(qs)}">Clear ticks</a>
    </div>
  </form>"""
        if not edit:
            form_html = (
                '<details class="more"><summary>Fix one row</summary>'
                + form_html + "</details>"
            )
    elif not can_write:
        form_html = ('<p class="sub">Viewing only — Managers and Team Leaders assign work. '
                     "You can still export the reports.</p>")

    load_plan_html = ""
    add_same_href = ""
    if can_write and project and not edit:
        keep_start = (q.get("plan_start", [""])[0] or "").strip()
        try:
            dt.date.fromisoformat(keep_start)
        except ValueError:
            keep_start = ""
        adding = (q.get("add", [""])[0] or "").strip() in ("1", "yes")
        saved = saved_plan_fields(db_path, project, start, days)
        q_floors = floor_days_from_q(q)
        if saved and adding:
            plan_start = keep_start or saved["start"]
            floor_days = q_floors
            load_plan_html = _loading_plan_html(
                ref, project, plan_start, hours_per_day,
                session_start=start,
                floor_days=floor_days,
                fill_days=(q.get("fill_floor_days", [""])[0] or "").strip(),
                zone_id=saved["building_id"],
                replace=False,
                append=True,
                week_qs=qs,
            )
        elif saved:
            plan_start = keep_start or saved["start"]
            floor_days = q_floors or saved["floor_days"]
            add_same_href = f"/forecast?{qs}&add=1"
            load_plan_html = _loading_plan_html(
                ref, project, plan_start, hours_per_day,
                session_start=start,
                floor_days=floor_days,
                fill_days=(q.get("fill_floor_days", [""])[0] or "").strip(),
                checked_people=saved["people"],
                checked_cats=saved["category_ids"],
                model_id=saved["model_id"],
                zone_id=saved["building_id"],
                replace=True,
                week_qs=qs,
            )
        else:
            plan_start = keep_start or next_plan_start(db_path, project, start, days)
            load_plan_html = _loading_plan_html(
                ref, project, plan_start, hours_per_day,
                session_start=start,
                floor_days=q_floors,
                fill_days=(q.get("fill_floor_days", [""])[0] or "").strip(),
            )

    exports = ""
    extra_exports = ""
    if project:
        fallback = ""
        if (exp_start, exp_end) != (start, end) and not rows:
            fallback = (
                f' <span class="hint">Saved plan {html.escape(exp_start)} → '
                f'{html.escape(exp_end)}</span>'
            )
        exports = (
            f'<a class="btn green" href="/forecast/export?{load_qs}&kind=loading">'
            f'Export this plan</a>{fallback}'
        )
        extra_exports = (
            f'<details class="more"><summary>Other Excel files</summary>'
            f'<div class="exports" style="margin-top:10px">'
            f'<a class="btn grey" href="/forecast/export?{qs}&kind=modelling">Modelling Works</a>'
            f'<a class="btn grey" href="/forecast/export?{qs}&kind=sheet">Sheet Work</a>'
            f'</div></details>'
        )

    grid = _load_grid_html(ref, project, dates, rows, hours_per_day, can_write, person_id) if project else ""
    table = _assignment_table(rows, show_zone, can_write, qs) if project else ""
    periods = saved_periods(db_path, project) if project else []
    periods_html = _periods_html(periods, project, start, days, can_write, add_same_href) if project else ""
    nxt = next_plan_start(db_path, project, start, days) if project else start
    next_qs = window_qs(project, nxt, days) if project else ""
    prog = project_progress(db_path, project) if project else {}
    progress_html = _progress_html(prog, qs, project, can_write) if project else ""
    if progress_html:
        progress_html = (
            '<details class="more"><summary>Project % complete</summary>'
            + progress_html + "</details>"
        )
    edit_json = json.dumps(edit) if edit else "null"
    ref_json = json.dumps(ref_page).replace("<", "\\u003c")
    top_form = form_html if edit else ""
    more_form = "" if edit else form_html

    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Forecast</title>{STYLE}</head><body><div class="wrap">
  <div class="nav"><a href="/">Task Entry</a> &nbsp;·&nbsp; <b>Forecast</b>
    &nbsp;·&nbsp; <a href="/admin">Manage Lists</a>
    &nbsp;·&nbsp; <a href="/reports">Reports</a></div>
  {banner}
  <h1>Forecast</h1>
  <p class="sub">Add floors for this week, then export. Faster than filling the loading sheet in Excel.
     Same week stays one plan. <b>New week</b> starts a new Excel.</p>
  {ok}
  {top_form}
  <form class="card window" method="GET" action="/forecast" id="window">
    <div><label>Project</label><select name="project" id="project" onchange="this.form.submit()">{proj_opts}</select></div>
    <div><label>This week from</label><input type="date" name="start" value="{html.escape(start)}"></div>
    <div><label>To</label><input type="date" name="end" value="{html.escape(end)}"></div>
    <button type="submit">Show</button>
    <a class="btn grey" href="/forecast?{html.escape(next_qs)}">New week</a>
  </form>
  {periods_html}
  {load_plan_html}
  <div class="section">
    <h2>This plan <span class="hint">{html.escape(start)} to {html.escape(window_end(start, days) if dates else "")}</span></h2>
    {exports}
  </div>
  {grid}
  {table}
  {extra_exports}
  {more_form}
  {progress_html}
</div>
<script>
const REF = {ref_json};
const EDIT = {edit_json};
const PROJECT = {json.dumps(project)};
const TEAM_SET = {json.dumps(bool(team_set))};
const FIRST_DAY = {json.dumps(start)};
const DAY_HRS = {HOURS_PER_DAY};
let hoursDirty = false;
function serviceVal(){{
  const e = document.querySelector('input[name=service]:checked');
  return e ? e.value : "Modelling Works";
}}
function ticked(id){{
  const box = document.getElementById(id); if(!box) return [];
  return [].slice.call(box.querySelectorAll("input:checked")).map(c=>c.value);
}}
function fillTicks(id, name, groups, keep){{
  const box = document.getElementById(id); if(!box) return;
  const on = new Set((keep || ticked(id)).map(String));
  let html = "";
  groups.forEach(g=>{{
    if(!g.items.length) return;
    if(g.label) html += `<div class="grp">${{g.label}}</div>`;
    html += g.items.map(i=>`<label><input type="checkbox" name="${{name}}" value="${{i.id}}"`
      + (on.has(String(i.id)) ? " checked" : "") + `>${{i.name}}</label>`).join("");
  }});
  box.innerHTML = html || `<span class="none">Nothing in this list yet.</span>`;
}}
function personLabel(p){{
  return p.emp_code ? (p.name + " (" + p.emp_code + ")") : p.name;
}}
function todayISO(){{
  const d=new Date();
  const m=d.getMonth()+1, day=d.getDate();
  return d.getFullYear()+"-"+(m<10?"0":"")+m+"-"+(day<10?"0":"")+day;
}}
function onTeamNow(row){{
  if(!row || row.project_code!=PROJECT) return false;
  if(String(row.left_on||"").trim()) return false;
  if(row.joined_on && row.joined_on > todayISO()) return false;
  return true;
}}
function currentMembership(pid){{
  return (REF.project_people||[]).find(x=>
    onTeamNow(x) && String(x.person_id)===String(pid)) || null;
}}
function fillPeople(){{
  const team=[...new Set((REF.project_people||[]).filter(onTeamNow).map(x=>String(x.person_id)))];
  const list = TEAM_SET ? REF.people.filter(p=>team.includes(String(p.id))) : REF.people;
  const proj = REF.projects.find(x=>x.code==PROJECT);
  const teamed = !!(proj && Number(proj.use_teams));
  const item = p=>({{id:p.id, name: personLabel(p)}});
  if(teamed){{
    const isLead = p=>{{
      const row = currentMembership(p.id);
      const role = (row && row.role) ? row.role : (p.role||"");
      return role.toUpperCase()==="QC REVIEWER";
    }};
    const isSupp = p=>{{
      const row = currentMembership(p.id);
      const role = (row && row.role) ? row.role : (p.role||"");
      return role.toUpperCase()==="QC SUPPORT";
    }};
    const leads = list.filter(isLead);
    const supps = list.filter(isSupp);
    const map = (REF.lead_support||[]).filter(r=>r.project_code==PROJECT);
    const groups = [{{label:"Team lead (QC Reviewer)", items: leads.map(item)}}];
    if(map.length){{
      leads.forEach(lead=>{{
        const sids = map.filter(r=>String(r.lead_id)===String(lead.id)).map(r=>String(r.support_id));
        const items = supps.filter(p=>sids.includes(String(p.id))).map(item);
        if(items.length) groups.push({{label:"QC Support → "+lead.name, items}});
      }});
      const assigned = new Set(map.map(r=>String(r.support_id)));
      const rest = supps.filter(p=>!assigned.has(String(p.id))).map(item);
      if(rest.length) groups.push({{label:"QC Support (not mapped)", items: rest}});
    }} else {{
      groups.push({{label:"Team support (QC Support)", items: supps.map(item)}});
    }}
    fillTicks("people_ids", "person_ids", groups);
  }} else {{
    fillTicks("people_ids", "person_ids", [{{label:"", items: list.map(item)}}]);
  }}
}}
function isLeadPerson(id){{
  const row = currentMembership(id);
  if(row && row.role) return String(row.role).toUpperCase()==="QC REVIEWER";
  const p = REF.people.find(x=>String(x.id)===String(id));
  return !!(p && (p.role||"").toUpperCase()==="QC REVIEWER");
}}
function leadScopeIds(kind, leadId){{
  if(!leadId) return [];
  return (REF.lead_scope||[]).filter(r=>r.project_code==PROJECT
    && String(r.person_id)==String(leadId) && r.kind==kind).map(r=>String(r.item_id));
}}
function unionScopeIds(kind){{
  const proj = REF.projects.find(x=>x.code==PROJECT);
  if(!proj || !Number(proj.use_teams)) return null;
  const leads = ticked("people_ids").filter(isLeadPerson);
  if(!leads.length) return null;
  const any = (REF.lead_scope||[]).some(r=>r.project_code==PROJECT && r.kind==kind);
  if(!any) return null;
  const u = new Set();
  leads.forEach(id=>leadScopeIds(kind, id).forEach(x=>u.add(x)));
  return [...u];
}}
function filterByUnion(items, idKey, kind){{
  const ids = unionScopeIds(kind);
  if(!ids) return items;
  if(kind==="zone"){{
    const allIds = items.filter(i=>Number(i.is_all)).map(i=>String(i.id));
    if(ids.some(id=>allIds.indexOf(id)>=0)) return items;
  }}
  return items.filter(i=>ids.indexOf(String(i[idKey]))>=0);
}}
function fillZoneLevelModel(){{
  const project = REF.projects.find(x=>x.code==PROJECT);
  const zoneWrap = document.getElementById("zone_wrap");
  const noZone = !!(project && project.no_zone);
  if(zoneWrap) zoneWrap.hidden = noZone;
  fillTicks("building_ids", "building_ids",
    [{{label:"", items: filterByUnion(REF.buildings.filter(b=>b.project_code==PROJECT), "id", "zone")
      .map(b=>({{id:b.id, name:b.name}}))}}]);
  fillTicks("level_ids", "level_ids",
    [{{label:"", items: filterByUnion(REF.levels.filter(l=>l.project_code==PROJECT), "id", "level")
      .map(l=>({{id:l.id, name:l.label}}))}}]);
  fillTicks("model_ids", "model_ids",
    [{{label:"", items: REF.models.filter(m=>m.project_code==PROJECT && m.name).map(m=>({{id:m.id, name:m.name}}))}}]);
}}
function fillCategories(){{
  fillTicks("category_ids", "category_ids",
    [{{label:"", items: filterByUnion(REF.categories, "id", "category")}}]);
}}
function fillTasks(){{
  const cats = ticked("category_ids");
  const taskIds = [...new Set(REF.category_tasks.filter(x=>!cats.length || cats.includes(String(x.category_id)))
                  .map(x=>String(x.task_name_id)))];
  const items = taskIds.length ? REF.task_names.filter(t=>taskIds.includes(String(t.id))) : REF.task_names;
  fillTicks("task_name_ids", "task_name_ids", [{{label:"", items: items}}]);
  fillSubTasks();
}}
function fillSubTasks(){{
  const cats = ticked("category_ids");
  const tasks = ticked("task_name_ids");
  const box = document.getElementById("task_sub_task_ids");
  if(!cats.length){{
    if(box) box.innerHTML = '<span class="none">Tick a Category to see Sub Tasks.</span>';
    updateCount();
    return;
  }}
  const byId = Object.fromEntries(REF.categories.map(c=>[String(c.id), c.name]));
  const groups = cats.map(cid=>({{
    label: byId[cid] || "",
    items: REF.task_sub_tasks.filter(s=>String(s.category_id)===String(cid)
      && (!tasks.length || tasks.includes(String(s.task_name_id))))
  }}));
  const loose = REF.task_sub_tasks.filter(s=>s.category_id==null
    && (!tasks.length || tasks.includes(String(s.task_name_id))));
  if(loose.length) groups.push({{label:"", items: loose}});
  fillTicks("task_sub_task_ids", "task_sub_task_ids", groups);
  updateCount();
}}
function toggleService(){{
  const modelling = serviceVal() === "Modelling Works";
  const m = document.getElementById("modelling_fields");
  if(m) m.hidden = !modelling;
  ["sheet_fields_name","sheet_fields_num","sheet_fields_scale","sheet_fields_paper"].forEach(id=>{{
    const e = document.getElementById(id); if(e) e.hidden = modelling;
  }});
  updateCount();
}}
function nTicks(id){{
  const e = document.getElementById(id); if(!e) return 0;
  return e.querySelectorAll("input[type=checkbox]:checked").length;
}}
function updateCount(){{
  const bar = document.getElementById("combo_count"); if(!bar) return;
  const project = REF.projects.find(x=>x.code==PROJECT);
  const noZone = !!(project && project.no_zone);
  const people = nTicks("people_ids");
  const dates = nTicks("work_dates");
  const floors = nTicks("level_ids");
  const missing = [];
  if(!people) missing.push("people");
  if(!dates) missing.push("days");
  if(!noZone && !nTicks("building_ids")) missing.push("zone");
  if(!floors) missing.push("floor");
  if(missing.length){{
    bar.innerHTML = "Still need: <b>" + missing.join(", ") + "</b>";
    return;
  }}
  if(floors > 1){{
    bar.innerHTML = "<b>" + people + " people × " + floors + " floors</b> is a loading plan. "
      + "Tick <b>one floor</b> here, or use <b>Add floors</b> above "
      + "(" + people + " × " + floors + " floors × 1 day = " + (people*floors)
      + " rows — not models × sub-tasks).";
    return;
  }}
  const total = people * dates;
  const extra = total > 250 ? " — tick fewer days (max 250 person-days here)." : "";
  const hrsEl = document.getElementById("estimated_hours");
  const personDays = people * dates;
  const maxHrs = Math.round(DAY_HRS * personDays * 100) / 100;
  if(hrsEl && !EDIT && !hoursDirty && personDays){{
    hrsEl.value = String(maxHrs);
  }}
  const totalHrs = hrsEl ? parseFloat(hrsEl.value) : 0;
  let hrsNote = "";
  if(totalHrs > 0 && personDays){{
    const perPersonDay = totalHrs / personDays;
    hrsNote = " <b>" + totalHrs + " hrs</b> for " + people + " person(s) × " + dates
      + " day(s) = <b>" + perPersonDay.toFixed(2) + " hrs per person per day</b>"
      + " (full day = " + DAY_HRS + " hrs).";
    if(totalHrs > maxHrs + 0.001){{
      hrsNote += " Cap is <b>" + maxHrs + " hrs</b>.";
    }}
  }}
  bar.innerHTML = "This will create <b>" + total + "</b> assignment" + (total===1?"":"s")
    + " (" + people + " people × " + dates + " days)." + hrsNote + extra;
}}
function tickVal(id, v){{
  if(v==null || v==="") return;
  const e = document.querySelector("#"+id+' input[value="'+String(v)+'"]');
  if(e) e.checked = true;
}}
function fill(){{
  fillPeople();
  fillZoneLevelModel();
  fillCategories();
  fillTasks();
  toggleService();
  updateCount();
}}
fill();
if(EDIT){{
  const set = (id,v)=>{{ const e=document.getElementById(id); if(e && v!=null && v!=="") e.value=String(v); }};
  document.querySelectorAll('input[name=service]').forEach(r=>{{ r.checked = r.value===EDIT.service; }});
  toggleService();
  tickVal("work_dates", EDIT.work_date);
  tickVal("people_ids", EDIT.person_id);
  fillZoneLevelModel();
  fillCategories();
  set("estimated_hours", EDIT.estimated_hours);
  tickVal("category_ids", EDIT.category_id);
  fillTasks();
  tickVal("task_name_ids", EDIT.task_name_id);
  fillSubTasks();
  tickVal("task_sub_task_ids", EDIT.task_sub_task_id);
  tickVal("building_ids", EDIT.building_id);
  tickVal("level_ids", EDIT.level_id);
  tickVal("model_ids", EDIT.model_id);
  set("area", EDIT.area);
  set("drawing_name", EDIT.drawing_name);
  set("drawing_number", EDIT.drawing_number);
  set("scale", EDIT.scale);
  set("paper_size", EDIT.paper_size);
  set("notes", EDIT.notes);
  updateCount();
}}
document.addEventListener("click", function(ev){{
  const all = ev.target.closest ? ev.target.closest("[data-tick-all],[data-tick-none]") : null;
  if(all){{
    ev.preventDefault();
    const on = all.hasAttribute("data-tick-all");
    const id = all.getAttribute(on ? "data-tick-all" : "data-tick-none");
    const box = document.getElementById(id);
    if(box) [].forEach.call(box.querySelectorAll("input[type=checkbox]"), c=>{{
      if(c.closest("label") && c.closest("label").style.display==="none") return;
      c.checked = on;
    }});
    if(id==="people_ids"){{ fillZoneLevelModel(); fillCategories(); fillTasks(); }}
    else if(id==="category_ids") fillTasks();
    else if(id==="task_name_ids") fillSubTasks();
    else if(id==="load_people" || id==="load_cats"){{
      var lp=document.getElementById("loading_plan");
      if(lp) lp.dispatchEvent(new Event("change"));
    }}
    else updateCount();
    return;
  }}
  const cell = ev.target.closest ? ev.target.closest("[data-date][data-person]") : null;
  if(!cell) return;
  tickVal("work_dates", cell.getAttribute("data-date"));
  tickVal("people_ids", cell.getAttribute("data-person"));
  updateCount();
  const form = document.getElementById("assign");
  if(form) form.scrollIntoView({{behavior:"smooth", block:"start"}});
}});
document.addEventListener("change", function(ev){{
  const t = ev.target;
  if(!t) return;
  if(t.name==="service"){{ toggleService(); return; }}
  const box = t.closest ? t.closest(".ticks, .days") : null;
  if(!box) return;
  if(box.id==="people_ids"){{ fillZoneLevelModel(); fillCategories(); fillTasks(); }}
  else if(box.id==="category_ids") fillTasks();
  else if(box.id==="task_name_ids") fillSubTasks();
  else updateCount();
}});
document.getElementById("assign") && document.getElementById("assign").addEventListener("submit", function(ev){{
  const people = nTicks("people_ids");
  const dates = nTicks("work_dates");
  const hrsEl = document.getElementById("estimated_hours");
  const totalHrs = hrsEl ? parseFloat(hrsEl.value) : 0;
  const maxHrs = DAY_HRS * people * dates;
  if(!EDIT && people && dates && totalHrs > maxHrs + 0.001){{
    ev.preventDefault();
    alert("Too many hours. " + people + " person(s) × " + dates + " day(s) can be at most "
      + maxHrs.toFixed(2) + " hrs (a day is " + DAY_HRS + " hrs: 09:00–18:30 minus 75 min breaks). "
      + "Sub-tasks share that day — they do not each get " + DAY_HRS + " hrs.");
  }}
}});
document.addEventListener("input", function(ev){{
  if(ev.target && ev.target.id==="estimated_hours"){{ hoursDirty = true; updateCount(); return; }}
  const id = ev.target.getAttribute && ev.target.getAttribute("data-filter");
  if(!id) return;
  const q = ev.target.value.toLowerCase();
  document.querySelectorAll("#"+id+" label").forEach(lab=>{{
    lab.style.display = lab.textContent.toLowerCase().indexOf(q)>=0 ? "" : "none";
  }});
}});
</script></body></html>"""
