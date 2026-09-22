#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Sync the Excel *Lists* sheet into the database WITHOUT touching tasks.

Use this for routine list maintenance: keep editing the Lists sheet in Excel the
way you always have (add a project, a building, a model, a new master task / MIDP
sheet under its sheet type, a person, etc.), then run this to push those additions
into the database. The team's entry form picks them up on the next page refresh.

  * ADDITIVE only — it inserts new list rows; it never deletes or renames, so it
    can't break tasks that already point at a building / level / master task.
  * Does NOT rebuild the database and NEVER touches tasks or assignees.
    (For a one-time full migration of the Master Log, use import_xlsx.py instead.)

Usage:
    python sync_lists.py                          # default workbook + tasklog.db
    python sync_lists.py "path/to.xlsx" tasklog.db
"""
import sys, os, re, sqlite3

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = r"C:\Users\jse2084\Documents\Copy of 2630_M45_Master Task Log.xlsx"
DEFAULT_DB = os.path.join(HERE, "tasklog.db")
DEFAULT_BREAKS = [("Lunch", "12:45", "13:30"), ("AM", "11:15", "11:30"), ("PM", "16:15", "16:30")]


def proj_code(v):
    if v is None:
        return None
    return ("%g" % v) if isinstance(v, (int, float)) else str(v).strip()


def main(xlsx_path, db_path):
    import openpyxl
    if not os.path.exists(db_path):
        sys.exit(f"Database not found: {db_path}. Run import_xlsx.py once first.")
    if not os.path.exists(xlsx_path):
        sys.exit(f"Workbook not found: {xlsx_path}")

    lists = openpyxl.load_workbook(xlsx_path, data_only=True)["Lists"]
    readme = openpyxl.load_workbook(xlsx_path, data_only=True)["Read Me"]
    con = sqlite3.connect(db_path)
    con.execute("PRAGMA foreign_keys = ON")
    cur = con.cursor()

    tables = ["projects", "break_windows", "buildings", "models", "master_tasks",
              "sheet_types", "levels", "task_names", "people", "statuses"]
    before = {t: cur.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in tables}

    def col_values(letter):
        c = openpyxl.utils.column_index_from_string(letter)
        return [lists.cell(r, c).value for r in range(2, lists.max_row + 1)
                if lists.cell(r, c).value not in (None, "")]

    hmap = {str(lists.cell(1, c).value).strip(): c
            for c in range(1, lists.max_column + 1)
            if lists.cell(1, c).value not in (None, "")}

    # ---- projects + break windows (names from Read Me) ----
    proj_names = {}
    for r in range(1, readme.max_row + 1):
        cells = [readme.cell(r, c).value for c in range(1, readme.max_column + 1)]
        cells = [str(c).strip() for c in cells if c not in (None, "")]
        if len(cells) >= 2 and re.fullmatch(r"\d+\.\d+", cells[0]):
            proj_names[cells[0]] = cells[1]
    for v in col_values("A"):
        code = proj_code(v)
        cur.execute("INSERT OR IGNORE INTO projects(code,name,client_code,internal_code) "
                    "VALUES (?,?,?,?)", (code, proj_names.get(code, f"Project {code}"), "M45", "2630"))
        for (bn, bs, be) in DEFAULT_BREAKS:
            cur.execute("INSERT OR IGNORE INTO break_windows(project_code,name,start_time,end_time) "
                        "VALUES (?,?,?,?)", (code, bn, bs, be))

    # ---- shared simple lists ----
    # Levels are per-project (Manage Lists / import). Do not seed a shared list.
    if hmap.get("Sheet Types"):
        for r in range(2, lists.max_row + 1):
            v = lists.cell(r, hmap["Sheet Types"]).value
            if v not in (None, ""):
                cur.execute("INSERT OR IGNORE INTO sheet_types(name) VALUES (?)", (str(v).strip(),))
    for v in col_values("I"):
        cur.execute("INSERT OR IGNORE INTO statuses(name) VALUES (?)", (str(v).strip(),))
    for v in col_values("H"):
        cur.execute("INSERT OR IGNORE INTO people(name) VALUES (?)", (str(v).strip(),))
    for letter in ("G", "W"):
        for v in col_values(letter):
            cur.execute("INSERT OR IGNORE INTO task_names(name) VALUES (?)", (str(v).strip(),))

    # ---- per-project buildings / models ----
    for hdr, c in hmap.items():
        mb = re.fullmatch(r"Building \((\d+\.\d+)\)", hdr)
        mm = re.fullmatch(r"Model Name \((\d+\.\d+)\)", hdr)
        if mb:
            code = mb.group(1)
            for r in range(2, lists.max_row + 1):
                v = lists.cell(r, c).value
                if v not in (None, ""):
                    nm = str(v).strip()
                    cur.execute("INSERT OR IGNORE INTO buildings(project_code,name,is_all) "
                                "VALUES (?,?,?)", (code, nm, 1 if nm.lower().startswith("all") else 0))
            cur.execute("INSERT OR IGNORE INTO buildings(project_code,name,is_all) VALUES (?,?,1)",
                        (code, "All Blocks"))
        elif mm:
            code = mm.group(1)
            for r in range(2, lists.max_row + 1):
                v = lists.cell(r, c).value
                if v not in (None, ""):
                    cur.execute("INSERT OR IGNORE INTO models(project_code,name) VALUES (?,?)",
                                (code, str(v).strip()))

    # ---- master tasks: columns headed '<proj>_<SheetType>' ----
    def sheet_type_id(name):
        row = cur.execute("SELECT id FROM sheet_types WHERE name=?", (name,)).fetchone()
        if row:
            return row[0]
        cur.execute("INSERT INTO sheet_types(name) VALUES (?)", (name,))
        return cur.lastrowid

    for hdr, c in hmap.items():
        m = re.fullmatch(r"(\d+\.\d+)_(.+)", hdr)
        if not m:
            continue
        code, stid = m.group(1), sheet_type_id(m.group(2).strip())
        for r in range(2, lists.max_row + 1):
            v = lists.cell(r, c).value
            if v in (None, ""):
                continue
            nm = re.sub(r"^\[\d+\]\s*", "", str(v).strip())
            cur.execute("INSERT OR IGNORE INTO master_tasks(project_code,sheet_type_id,name) "
                        "VALUES (?,?,?)", (code, stid, nm))

    con.commit()
    after = {t: cur.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in tables}
    tasks_n = cur.execute("SELECT COUNT(*) FROM tasks").fetchone()[0]
    con.close()

    print(f"=== LIST SYNC COMPLETE -> {db_path} ===")
    any_change = False
    for t in tables:
        delta = after[t] - before[t]
        flag = f"  (+{delta} new)" if delta else ""
        if delta:
            any_change = True
        print(f"  {t:14} {after[t]:>4}{flag}")
    print(f"\n  tasks untouched: {tasks_n}")
    print("  No list changes found." if not any_change
          else "  New list items are now available in the entry form (refresh the page).")


if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    db = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_DB
    main(xlsx, db)
